"""
Optimizador Black-Litterman — aplicación Streamlit autocontenida (un solo archivo).

Modelo Black-Litterman de nivel profesional con datos de Yahoo Finance y moneda base USD:
equilibrio de mercado (Π = δΣw_mkt), views absolutas/relativas/de canasta con Ω de
He-Litterman, Idzorek o intervalos, distribución posterior de He-Litterman (1999),
optimización media-varianza con restricciones (cvxpy), métricas de riesgo, backtest
walk-forward y exportación a Excel, JSON y PDF.

Ejecutar:  streamlit run app.py
"""
from __future__ import annotations


import hashlib
import importlib.util
import io
import json
import os
import pickle
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime, timedelta

import streamlit as st

# Diagnóstico claro si falta alguna dependencia de requirements.txt
_MISSING = [m for m in ("numpy", "pandas", "scipy", "cvxpy", "plotly", "yfinance", "openpyxl",
                        "reportlab", "matplotlib") if importlib.util.find_spec(m) is None]
if _MISSING:
    st.set_page_config(page_title="Black-Litterman Optimizer", page_icon="📈")
    st.error("Faltan paquetes por instalar. Verifica que requirements.txt esté en la raíz del repositorio "
             "(junto a app.py) y reinicia la app.")
    st.code("\n".join(_MISSING))
    st.stop()

import cvxpy as cp  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import plotly.graph_objects as go  # noqa: E402
from scipy import optimize, stats  # noqa: E402

__version__ = "1.1.0"


# ==============================================================================
# CONFIGURACIÓN Y CONSTANTES
# ==============================================================================
# Parámetros por defecto y constantes del modelo.
TRADING_DAYS = 252

# --- Opciones (etiquetas mostradas en la app → clave interna) -----------------
WMKT_METHODS = {"Capitalización de mercado": "market_cap",
                "Pesos iguales (1/N)": "equal",
                "Pesos manuales": "manual"}
DELTA_METHODS = {"Fijo estándar (2.5)": "fixed",
                 "Implícito del mercado": "implied",
                 "Manual": "manual"}
COV_METHODS = {"Ledoit-Wolf (correlación constante)": "ledoit_wolf",
               "Muestral": "sample",
               "EWMA (RiskMetrics)": "ewma"}
TAU_METHODS = {"Estándar (0.05)": "fixed",
               "1/T (años de datos)": "inverse_t",
               "Manual": "manual"}
OMEGA_METHODS = {"Idzorek (confianza %)": "idzorek",
                 "He-Litterman (proporcional a la varianza)": "he_litterman",
                 "Intervalos de confianza": "interval"}
BUDGET_MODES = {"Totalmente invertido (Σw = 1)": "full",
                "Sin apalancamiento (Σw ≤ 1, resto en T-Bills)": "no_leverage",
                "Libre (permite apalancamiento)": "free"}
SIGMA_OPT = {"Σ_BL posterior (He-Litterman)": "posterior",
             "Σ histórica (prior)": "prior"}
RF_METHODS = {"^IRX último dato (T-Bill 13 semanas)": "irx_last",
              "^IRX promedio del periodo": "irx_mean",
              "Manual": "manual"}

VIEW_TYPES = ["Absoluta", "Relativa", "Canasta"]
BASKET_WEIGHTING = ["Igual", "Capitalización"]

MIN_ASSETS_OPT = 2
MAX_ASSETS = 10


@dataclass
class ModelConfig:
    """Configuración completa del modelo (serializable a JSON)."""
    tickers: list[str] = field(default_factory=lambda: [
        "AAPL", "MSFT", "NVDA", "JPM", "XOM", "JNJ", "WALMEX.MX", "GMEXICOB.MX"])
    start: str = field(default_factory=lambda: (date.today() - timedelta(days=5 * 365 + 1)).isoformat())
    end: str = field(default_factory=lambda: date.today().isoformat())
    benchmark: str = "^GSPC"
    max_missing_pct: float = 10.0

    rf_method: str = "irx_last"
    rf_manual_pct: float = 4.0

    wmkt_method: str = "market_cap"
    manual_weights: dict[str, float] = field(default_factory=dict)

    delta_method: str = "fixed"
    delta_manual: float = 2.5
    delta_bounds: tuple[float, float] = (1.0, 10.0)

    cov_method: str = "ledoit_wolf"
    ewma_lambda: float = 0.94

    tau_method: str = "fixed"
    tau_manual: float = 0.05

    omega_method: str = "idzorek"
    interval_prob_pct: float = 90.0

    sigma_opt: str = "posterior"
    budget: str = "full"
    long_only: bool = True
    use_asset_bounds: bool = True
    default_min_pct: float = 0.0
    default_max_pct: float = 40.0
    asset_bounds: dict[str, list[float]] = field(default_factory=dict)   # ticker → [min%, max%]
    use_groups: bool = False
    asset_groups: dict[str, str] = field(default_factory=dict)          # ticker → grupo
    group_bounds: dict[str, list[float]] = field(default_factory=dict)  # grupo → [min%, max%]
    use_turnover: bool = False
    turnover_max_pct: float = 30.0
    current_weights: dict[str, float] = field(default_factory=dict)     # ticker → %

    views: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        d = asdict(self)
        d["delta_bounds"] = list(self.delta_bounds)
        return d

    @classmethod
    def from_dict(cls, d: dict) -> "ModelConfig":
        known = {k: v for k, v in d.items() if k in cls.__dataclass_fields__}
        cfg = cls(**known)
        cfg.delta_bounds = tuple(cfg.delta_bounds)
        return cfg


# ==============================================================================
# DATOS (YAHOO FINANCE → USD)
# ==============================================================================
# Capa de datos: descarga desde Yahoo Finance, limpieza, alineación de
# calendarios y conversión de todos los precios a USD.
#
# Convenciones
# ------------
# * Precios de cierre ajustados (``auto_adjust=True``): incorporan dividendos y
#   splits, por lo que los rendimientos son rendimientos totales.
# * Tipo de cambio: el ticker ``"{CCY}=X"`` de Yahoo cotiza *unidades de CCY por
#   1 USD*; por tanto ``precio_USD = precio_local / FX``.
# * Monedas cotizadas en subunidades (GBp, ILA, ZAc) se dividen entre 100.
# * Tasa libre de riesgo: ^IRX (T-Bill 13 semanas), expresado en % anual.
#
# Para pruebas sin conexión, la variable de entorno ``BL_DEMO_DATA=1`` sustituye
# Yahoo Finance por un generador sintético reproducible.
SUBUNIT_CURRENCIES = {"GBp": ("GBP", 100.0), "GBX": ("GBP", 100.0),
                      "ILA": ("ILS", 100.0), "ZAc": ("ZAR", 100.0)}
RF_TICKER = "^IRX"
FFILL_LIMIT = 5  # días hábiles máximos a rellenar (feriados locales)


@dataclass
class MarketData:
    prices_usd: pd.DataFrame            # precios de los activos en USD
    returns: pd.DataFrame               # rendimientos simples diarios (USD)
    benchmark_prices: pd.Series
    benchmark_returns: pd.Series
    rf_daily: pd.Series                 # tasa libre de riesgo diaria (decimal)
    rf_annual_last: float               # último ^IRX (decimal anual)
    rf_annual_mean: float               # promedio ^IRX del periodo (decimal anual)
    currencies: dict[str, str]
    market_caps_usd: dict[str, float | None]
    names: dict[str, str]
    dropped: dict[str, str] = field(default_factory=dict)   # ticker → motivo
    warnings: list[str] = field(default_factory=list)
    missing_pct: dict[str, float] = field(default_factory=dict)

    @property
    def tickers(self) -> list[str]:
        return list(self.prices_usd.columns)

    @property
    def n_obs(self) -> int:
        return len(self.returns)

    @property
    def years(self) -> float:
        return self.n_obs / TRADING_DAYS


# =============================================================================
# Proveedores
# =============================================================================
def _extract_close(raw: pd.DataFrame, tickers: list[str]) -> pd.DataFrame:
    """Extrae la matriz de cierres de la salida de ``yf.download`` (robusto a versiones)."""
    if raw is None or raw.empty:
        return pd.DataFrame(columns=tickers, dtype=float)
    if isinstance(raw.columns, pd.MultiIndex):
        lvl0 = raw.columns.get_level_values(0)
        if "Close" in lvl0:
            close = raw["Close"]
        elif "Adj Close" in lvl0:
            close = raw["Adj Close"]
        else:  # group_by="ticker"
            close = raw.xs("Close", axis=1, level=1)
    else:
        close = raw[["Close"]].rename(columns={"Close": tickers[0]}) if "Close" in raw else raw
    if isinstance(close, pd.Series):
        close = close.to_frame(tickers[0])
    close = close.reindex(columns=tickers)
    close.index = pd.to_datetime(close.index).tz_localize(None).normalize()
    return close.astype(float)


class YahooProvider:
    """Proveedor de datos basado en ``yfinance``."""

    def download_prices(self, tickers: list[str], start: str, end: str, retries: int = 3) -> pd.DataFrame:
        """Descarga con reintentos y respaldo por ticker (Yahoo limita la tasa de consultas)."""
        import time
        import yfinance as yf
        end_inclusive = (pd.Timestamp(end) + pd.Timedelta(days=1)).strftime("%Y-%m-%d")
        close = pd.DataFrame(columns=tickers, dtype=float)
        for attempt in range(retries):
            try:
                raw = yf.download(tickers, start=start, end=end_inclusive, auto_adjust=True,
                                  progress=False, threads=False, group_by="column")
                close = _extract_close(raw, tickers)
            except Exception:  # noqa: BLE001
                close = pd.DataFrame(columns=tickers, dtype=float)
            missing = [t for t in tickers if t not in close or close[t].isna().all()]
            if not missing:
                return close
            time.sleep(1.5 * (attempt + 1))
        # respaldo: tickers faltantes uno por uno
        for t in [t for t in tickers if t not in close or close[t].isna().all()]:
            try:
                h = yf.Ticker(t).history(start=start, end=end_inclusive, auto_adjust=True)
                if not h.empty:
                    s = h["Close"].copy()
                    s.index = pd.to_datetime(s.index).tz_localize(None).normalize()
                    close = close.reindex(close.index.union(s.index)) if not close.empty else pd.DataFrame(index=s.index, columns=tickers, dtype=float)
                    close[t] = s
            except Exception:  # noqa: BLE001
                continue
        return close.reindex(columns=tickers).sort_index()

    def metadata(self, ticker: str) -> dict:
        """Moneda, capitalización (o activos totales para ETFs) y nombre."""
        import yfinance as yf
        t = yf.Ticker(ticker)
        out = {"currency": None, "market_cap": None, "name": ticker}
        try:
            fi = t.fast_info
            out["currency"] = fi.get("currency") if hasattr(fi, "get") else fi["currency"]
            mc = fi.get("market_cap") if hasattr(fi, "get") else fi["market_cap"]
            out["market_cap"] = float(mc) if mc else None
        except Exception:
            pass
        if out["currency"] is None or out["market_cap"] is None:
            try:
                info = t.info or {}
                out["currency"] = out["currency"] or info.get("currency")
                cap = info.get("marketCap") or info.get("totalAssets")
                out["market_cap"] = out["market_cap"] or (float(cap) if cap else None)
                out["name"] = info.get("shortName") or info.get("longName") or ticker
            except Exception:
                pass
        return out


class SyntheticProvider:
    """Datos sintéticos reproducibles (solo para pruebas: ``BL_DEMO_DATA=1``)."""

    def __init__(self, seed: int = 7):
        self.seed = seed

    def _calendar(self, start, end):
        return pd.bdate_range(start, end)

    def download_prices(self, tickers, start, end):
        idx = self._calendar(start, end)
        rng = np.random.default_rng(self.seed)
        n = len(tickers)
        out = {}
        # factor de mercado + factores idiosincráticos
        mkt = rng.normal(0.0004, 0.011, len(idx))
        for i, tk in enumerate(tickers):
            r_i = np.random.default_rng(self.seed + sum(map(ord, tk)))
            if tk == "^IRX":
                out[tk] = 4.5 + np.cumsum(r_i.normal(0, 0.02, len(idx))).clip(-4, 2)
                continue
            if tk.endswith("=X"):
                out[tk] = 18.0 * np.exp(np.cumsum(r_i.normal(0, 0.006, len(idx))))
                continue
            beta = 0.6 + 0.8 * r_i.random()
            idio = r_i.normal(0.0002 * r_i.random(), 0.012 + 0.01 * r_i.random(), len(idx))
            r = beta * mkt + idio if not tk.startswith("^") else mkt
            out[tk] = 100 * np.exp(np.cumsum(r))
        df = pd.DataFrame(out, index=idx)
        # feriados locales para emisoras mexicanas
        mx = [t for t in tickers if t.endswith(".MX")]
        if mx:
            holes = np.random.default_rng(self.seed).choice(len(idx), size=len(idx) // 40, replace=False)
            df.iloc[holes, [df.columns.get_loc(c) for c in mx]] = np.nan
        return df.reindex(columns=tickers)

    def metadata(self, ticker):
        ccy = "MXN" if ticker.endswith(".MX") else "USD"
        cap = 5e10 + 1e10 * (sum(map(ord, ticker)) % 40)
        return {"currency": ccy, "market_cap": cap, "name": f"{ticker} (sintético)"}


def get_provider():
    return SyntheticProvider() if os.environ.get("BL_DEMO_DATA") == "1" else YahooProvider()


# =============================================================================
# Pipeline
# =============================================================================
def parse_tickers(text: str) -> list[str]:
    """Normaliza la captura libre: separa por comas/espacios, mayúsculas, sin duplicados."""
    raw = [t.strip().upper() for t in text.replace(";", ",").replace("\n", ",").replace(" ", ",").split(",")]
    seen, out = set(), []
    for t in raw:
        if t and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def _fx_ticker(ccy: str) -> str:
    return f"{ccy}=X"


def load_market_data(tickers: list[str], start: str, end: str, benchmark: str = "^GSPC",
                     max_missing_pct: float = 10.0, provider=None) -> MarketData:
    """Descarga, limpia, alinea y convierte a USD el universo de inversión."""
    provider = provider or get_provider()
    warnings: list[str] = []
    dropped: dict[str, str] = {}

    # --- metadatos -----------------------------------------------------------
    meta = {t: provider.metadata(t) for t in tickers + [benchmark]}
    currencies, divisors = {}, {}
    for t, m in meta.items():
        ccy = m.get("currency") or "USD"
        if m.get("currency") is None:
            warnings.append(f"{t}: Yahoo no reportó la divisa; se asume USD.")
        base, div = SUBUNIT_CURRENCIES.get(ccy, (ccy.upper(), 1.0))
        currencies[t], divisors[t] = base, div

    fx_needed = sorted({c for c in currencies.values() if c != "USD"})
    fx_tickers = [_fx_ticker(c) for c in fx_needed]

    # --- precios ---------------------------------------------------------------
    all_tk = tickers + [benchmark, RF_TICKER] + fx_tickers
    px = provider.download_prices(all_tk, start, end)
    px = px[~px.index.duplicated(keep="last")].sort_index()
    # filas donde no cotizó ningún activo del universo
    px = px.loc[px[tickers + [benchmark]].notna().any(axis=1)]
    if px.empty:
        raise ValueError("Yahoo Finance no devolvió datos para el periodo solicitado.")

    # --- control de calidad -------------------------------------------------------
    missing_pct = {}
    for t in tickers:
        miss = float(px[t].isna().mean() * 100) if t in px else 100.0
        missing_pct[t] = miss
        if miss >= 99.9:
            dropped[t] = "sin datos (ticker inválido o sin historia)"
        elif miss > max_missing_pct:
            dropped[t] = f"{miss:.1f}% de datos faltantes (> {max_missing_pct:.0f}%)"
    for t in fx_tickers:
        if px[t].isna().all():
            for tk, c in currencies.items():
                if _fx_ticker(c) == t and tk in tickers and tk not in dropped:
                    dropped[tk] = f"sin tipo de cambio {t} para convertir a USD"
    if px[benchmark].isna().all():
        raise ValueError(f"No hay datos para el índice de referencia {benchmark}.")
    kept = [t for t in tickers if t not in dropped]
    if not kept:
        raise ValueError("Ningún ticker superó el control de calidad de datos.")

    # --- alineación de calendarios ---------------------------------------------------
    px = px.ffill(limit=FFILL_LIMIT)
    px = px.dropna(subset=kept + [benchmark])

    # --- conversión a USD -----------------------------------------------------------
    def to_usd(t):
        s = px[t] / divisors[t]
        c = currencies[t]
        return s if c == "USD" else s / px[_fx_ticker(c)]

    prices_usd = pd.DataFrame({t: to_usd(t) for t in kept}).dropna()
    bench = to_usd(benchmark).reindex(prices_usd.index).ffill()

    returns = prices_usd.pct_change().iloc[1:]
    bench_ret = bench.pct_change().reindex(returns.index)

    # --- tasa libre de riesgo ------------------------------------------------------
    irx = px[RF_TICKER].reindex(prices_usd.index).ffill().bfill() if RF_TICKER in px else None
    if irx is None or irx.isna().all():
        warnings.append("No se obtuvo ^IRX; se usa tasa libre de riesgo de 0%.")
        irx = pd.Series(0.0, index=prices_usd.index)
    rf_annual = irx / 100.0
    rf_daily = (rf_annual / TRADING_DAYS).reindex(returns.index)

    # --- capitalizaciones en USD ----------------------------------------------------------
    caps = {}
    for t in kept:
        mc = meta[t].get("market_cap")
        c = currencies[t]
        if mc is None:
            caps[t] = None
            continue
        mc = mc / divisors[t]
        if c != "USD":
            mc = mc / float(px[_fx_ticker(c)].dropna().iloc[-1])
        caps[t] = float(mc)

    if len(returns) < 60:
        warnings.append(f"Solo {len(returns)} observaciones: la estimación de Σ será poco confiable.")

    return MarketData(prices_usd=prices_usd, returns=returns, benchmark_prices=bench,
                      benchmark_returns=bench_ret, rf_daily=rf_daily,
                      rf_annual_last=float(rf_annual.iloc[-1]),
                      rf_annual_mean=float(rf_annual.mean()),
                      currencies={t: currencies[t] for t in kept},
                      market_caps_usd=caps,
                      names={t: meta[t].get("name", t) for t in kept},
                      dropped=dropped, warnings=warnings, missing_pct=missing_pct)


def annualized_stats(returns: pd.DataFrame) -> pd.DataFrame:
    """Estadísticos históricos anualizados (rendimiento aritmético, volatilidad)."""
    mu = returns.mean() * TRADING_DAYS
    vol = returns.std(ddof=1) * np.sqrt(TRADING_DAYS)
    cum = (1 + returns).prod()
    years = len(returns) / TRADING_DAYS
    cagr = cum ** (1 / years) - 1
    dd = ((1 + returns).cumprod() / (1 + returns).cumprod().cummax() - 1).min()
    return pd.DataFrame({"Rend. aritmético anual": mu, "CAGR": cagr,
                         "Volatilidad anual": vol, "Máx. drawdown": dd})


# ==============================================================================
# MATRIZ DE COVARIANZA Σ
# ==============================================================================
# Estimadores de la matriz de covarianza anualizada Σ.
#
# * Muestral: estimador insesgado clásico.
# * Ledoit-Wolf (2004), "Honey, I Shrunk the Sample Covariance Matrix":
#   Σ̂ = δ* F + (1 − δ*) S, con F la matriz de correlación constante y δ* la
#   intensidad de contracción óptima (asintóticamente) bajo pérdida de Frobenius.
# * EWMA (RiskMetrics, 1996): Σ_t = λ Σ_{t−1} + (1 − λ) r_t r_tᵀ, media cero.
@dataclass
class CovResult:
    cov: pd.DataFrame           # anualizada
    method: str
    shrinkage: float | None = None
    condition_number: float = float("nan")
    min_eigenvalue: float = float("nan")
    psd_repaired: bool = False


def sample_cov(returns: pd.DataFrame) -> np.ndarray:
    return np.cov(returns.values, rowvar=False, ddof=1)


def ledoit_wolf_constant_correlation(returns: pd.DataFrame) -> tuple[np.ndarray, float]:
    """Contracción hacia correlación constante (Ledoit & Wolf, 2004).

    Implementación fiel a ``covCor.m`` de los autores. Devuelve (Σ diaria, δ*).
    """
    x = returns.values
    t, n = x.shape
    if n == 1:
        return np.atleast_2d(np.var(x, ddof=1)), 0.0
    y = x - x.mean(axis=0)
    sample = y.T @ y / t
    var = np.diag(sample)
    sd = np.sqrt(var)
    corr = sample / np.outer(sd, sd)
    r_bar = (corr.sum() - n) / (n * (n - 1))
    prior = r_bar * np.outer(sd, sd)
    np.fill_diagonal(prior, var)

    # π̂: suma de varianzas asintóticas de las entradas de S
    y2 = y ** 2
    phi_mat = y2.T @ y2 / t - sample ** 2
    phi = phi_mat.sum()

    # ρ̂: suma de covarianzas asintóticas entre F y S
    term1 = (y ** 3).T @ y / t
    term3 = sample * var[:, None]
    theta_mat = term1 - term3
    np.fill_diagonal(theta_mat, 0.0)
    rho = np.trace(phi_mat) + r_bar * ((np.outer(1 / sd, sd)) * theta_mat).sum()

    # γ̂: error de especificación del objetivo
    gamma = np.linalg.norm(sample - prior, "fro") ** 2
    kappa = (phi - rho) / gamma if gamma > 0 else 0.0
    shrink = float(max(0.0, min(1.0, kappa / t)))
    sigma = shrink * prior + (1 - shrink) * sample
    return sigma * t / (t - 1), shrink   # reescala a denominador T−1 (consistente con la muestral)


def ewma_cov(returns: pd.DataFrame, lam: float = 0.94) -> np.ndarray:
    """EWMA RiskMetrics con media cero; pesos normalizados (suman 1)."""
    x = returns.values
    t = x.shape[0]
    w = (1 - lam) * lam ** np.arange(t - 1, -1, -1)
    w = w / w.sum()
    return (x * w[:, None]).T @ x


def nearest_psd(mat: np.ndarray, eps: float = 1e-12) -> tuple[np.ndarray, bool]:
    """Simetriza y recorta eigenvalores negativos (Higham simplificado)."""
    sym = (mat + mat.T) / 2
    vals, vecs = np.linalg.eigh(sym)
    if vals.min() >= eps:
        return sym, False
    vals = np.clip(vals, eps, None)
    return (vecs * vals) @ vecs.T, True


def estimate_covariance(returns: pd.DataFrame, method: str = "ledoit_wolf",
                        ewma_lambda: float = 0.94) -> CovResult:
    """Σ anualizada (× 252) con diagnóstico de condicionamiento."""
    shrink = None
    if method == "sample":
        daily = sample_cov(returns)
    elif method == "ledoit_wolf":
        daily, shrink = ledoit_wolf_constant_correlation(returns)
    elif method == "ewma":
        daily = ewma_cov(returns, ewma_lambda)
    else:
        raise ValueError(f"Método de covarianza desconocido: {method}")
    ann = np.atleast_2d(daily) * TRADING_DAYS
    ann, repaired = nearest_psd(ann)
    eig = np.linalg.eigvalsh(ann)
    cov = pd.DataFrame(ann, index=returns.columns, columns=returns.columns)
    return CovResult(cov=cov, method=method, shrinkage=shrink,
                     condition_number=float(eig.max() / eig.min()),
                     min_eigenvalue=float(eig.min()), psd_repaired=repaired)


def cov_to_corr(cov: pd.DataFrame) -> pd.DataFrame:
    sd = np.sqrt(np.diag(cov.values))
    return pd.DataFrame(cov.values / np.outer(sd, sd), index=cov.index, columns=cov.columns)


# ==============================================================================
# EQUILIBRIO DE MERCADO (Π, δ, τ)
# ==============================================================================
# Equilibrio de mercado (optimización inversa).
#
#     Π = δ Σ w_mkt           (rendimientos implícitos en exceso de r_f)
#
# δ implícito:  δ = (E[R_m] − r_f) / σ²_m, estimado con el índice de referencia.
DELTA_STANDARD = 2.5
TAU_STANDARD = 0.05


@dataclass
class EquilibriumResult:
    w_mkt: pd.Series
    wmkt_method_used: str
    delta: float
    delta_raw: float | None
    tau: float
    pi: pd.Series                      # exceso anual
    warnings: list[str] = field(default_factory=list)


def market_weights(tickers: list[str], method: str, market_caps: dict[str, float | None],
                   manual: dict[str, float] | None = None) -> tuple[pd.Series, str, list[str]]:
    """Pesos del portafolio de referencia. Devuelve (w, método usado, advertencias)."""
    warns: list[str] = []
    n = len(tickers)
    if method == "market_cap":
        missing = [t for t in tickers if not market_caps.get(t)]
        if missing:
            warns.append("Sin capitalización de mercado para " + ", ".join(missing)
                         + ": se usa 1/N como respaldo para todo el universo.")
            return pd.Series(1 / n, index=tickers), "equal", warns
        caps = pd.Series({t: market_caps[t] for t in tickers}, dtype=float)
        return caps / caps.sum(), "market_cap", warns
    if method == "manual":
        m = pd.Series({t: float((manual or {}).get(t, 0.0)) for t in tickers})
        if (m < 0).any():
            warns.append("Pesos manuales negativos: se interpretan como posiciones cortas del benchmark.")
        if abs(m.sum()) < 1e-12:
            warns.append("Pesos manuales suman 0: se usa 1/N.")
            return pd.Series(1 / n, index=tickers), "equal", warns
        if abs(m.sum() - 1) > 1e-6 and abs(m.sum() - 100) > 1e-6:
            warns.append(f"Los pesos manuales suman {m.sum():.2f}; se normalizan a 100%.")
        return m / m.sum(), "manual", warns
    return pd.Series(1 / n, index=tickers), "equal", warns


def implied_delta(bench_returns: pd.Series, rf_daily: pd.Series) -> float:
    """δ = E[R_m − r_f] / σ²_m (anualizados)."""
    ex = (bench_returns - rf_daily.reindex(bench_returns.index).fillna(0)).dropna()
    premium = ex.mean() * TRADING_DAYS
    var = bench_returns.dropna().var(ddof=1) * TRADING_DAYS
    return float(premium / var)


def resolve_delta(method: str, manual: float, bench_returns: pd.Series | None,
                  rf_daily: pd.Series | None, bounds=(1.0, 10.0)) -> tuple[float, float | None, list[str]]:
    warns: list[str] = []
    if method == "manual":
        return float(manual), None, warns
    if method == "implied":
        raw = implied_delta(bench_returns, rf_daily)
        if not np.isfinite(raw) or raw <= 0:
            warns.append(f"δ implícito = {raw:.2f} (prima de mercado no positiva en la muestra); "
                         f"se usa el valor estándar {DELTA_STANDARD}.")
            return DELTA_STANDARD, raw, warns
        lo, hi = bounds
        if raw < lo or raw > hi:
            clipped = float(np.clip(raw, lo, hi))
            warns.append(f"δ implícito = {raw:.2f} fuera de [{lo}, {hi}]; se acota a {clipped:.2f}.")
            return clipped, raw, warns
        return raw, raw, warns
    return DELTA_STANDARD, None, warns


def resolve_tau(method: str, manual: float, years: float) -> tuple[float, str]:
    """τ. Con 1/T se usa T en años: Var(μ̂_anual) = Σ_anual / T_años (error estándar de la media)."""
    if method == "inverse_t":
        tau = 1.0 / max(years, 1e-6)
        return float(min(tau, 1.0)), f"τ = 1/T con T = {years:.2f} años de datos"
    if method == "manual":
        return float(manual), "τ manual"
    return TAU_STANDARD, "τ estándar (Black-Litterman 1992; He-Litterman 1999)"


def implied_returns(cov: pd.DataFrame, w_mkt: pd.Series, delta: float) -> pd.Series:
    return pd.Series(delta * cov.values @ w_mkt.reindex(cov.index).values, index=cov.index)


# ==============================================================================
# VIEWS (P, Q, Ω)
# ==============================================================================
# Views del inversionista:  P μ = Q + ε,   ε ~ N(0, Ω).
#
# Tipos soportados
# ----------------
# * Absoluta : un activo;          "AAPL rendirá 12% anual"          → Q_exceso = Q − r_f
# * Relativa : un activo vs otro;  "MSFT superará a AAPL por 3%"     → Q (r_f se cancela)
# * Canasta  : n activos vs m;     "AAPL+MSFT superarán a JPM+BAC por 4%"
#              (pesos iguales o por capitalización dentro de cada lado;
#               sin lado corto se trata como view absoluta de la canasta).
#
# Métodos para Ω
# --------------
# * He-Litterman (1999):  ω_k = p_k (τΣ) p_kᵀ
# * Idzorek (2005):       confianza c_k ∈ (0, 1]; ω_k tal que el portafolio se
#   desplace exactamente la fracción c_k desde w_mkt hacia w_100%.
#   Solución cerrada (Walters, 2014): ω_k = ((1 − c_k)/c_k) · p_k (τΣ) p_kᵀ.
#   Se incluye también la calibración numérica original para verificación.
# * Intervalos:           [L, U] con probabilidad p → ω_k = ((U − L) / (2 z_{(1+p)/2}))²
CONF_MIN, CONF_MAX = 1e-4, 1.0 - 1e-6

VIEW_COLUMNS = ["Activa", "Tipo", "Activos largo", "Activos corto", "Rendimiento esperado (%)",
                "Confianza (%)", "Ponderación canasta", "Límite inferior (%)", "Límite superior (%)",
                "Descripción"]


@dataclass
class ViewSet:
    P: pd.DataFrame                  # k × n
    Q: pd.Series                     # exceso anual (decimal)
    Q_input: pd.Series               # tal como se capturó (decimal)
    confidence: pd.Series            # decimal
    lower: pd.Series
    upper: pd.Series
    labels: list[str]
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def k(self) -> int:
        return len(self.labels)


def empty_views_df() -> pd.DataFrame:
    return pd.DataFrame(columns=VIEW_COLUMNS).astype({
        "Activa": bool, "Tipo": str, "Activos largo": str, "Activos corto": str,
        "Rendimiento esperado (%)": float, "Confianza (%)": float, "Ponderación canasta": str,
        "Límite inferior (%)": float, "Límite superior (%)": float, "Descripción": str})


def example_views_df(tickers: list[str]) -> pd.DataFrame:
    """Views de ejemplo construidas con los primeros tickers del universo."""
    rows = []
    if len(tickers) >= 1:
        rows.append([True, "Absoluta", tickers[0], "", 12.0, 50.0, "Igual", 6.0, 18.0,
                     f"{tickers[0]} rendirá 12% anual"])
    if len(tickers) >= 3:
        rows.append([True, "Relativa", tickers[1], tickers[2], 3.0, 40.0, "Igual", 0.0, 6.0,
                     f"{tickers[1]} superará a {tickers[2]} por 3%"])
    if len(tickers) >= 6:
        rows.append([True, "Canasta", f"{tickers[3]}, {tickers[4]}", f"{tickers[5]}", 2.0, 25.0,
                     "Capitalización", -2.0, 6.0, "Canasta larga vs canasta corta"])
    df = pd.DataFrame(rows, columns=VIEW_COLUMNS)
    return df if not df.empty else empty_views_df()


def _split(cell) -> list[str]:
    if cell is None or (isinstance(cell, float) and np.isnan(cell)):
        return []
    return [s.strip().upper() for s in str(cell).replace(";", ",").split(",") if s.strip()]


def _num(x, default=np.nan) -> float:
    try:
        v = float(x)
        return default if np.isnan(v) else v
    except (TypeError, ValueError):
        return default


def build_views(df: pd.DataFrame, tickers: list[str], rf: float,
                market_caps: dict[str, float | None] | None = None) -> ViewSet:
    """Traduce la tabla de views a (P, Q). Q en exceso de r_f (decimal anual)."""
    errors, warns = [], []
    rows_p, q_ex, q_in, conf, lo, up, labels = [], [], [], [], [], [], []
    caps = market_caps or {}
    if df is None:
        df = empty_views_df()
    for i, r in df.reset_index(drop=True).iterrows():
        tag = f"View {i + 1}"
        active = r.get("Activa", True)
        if active is False or (isinstance(active, float) and np.isnan(active)) or str(active).lower() == "false":
            continue
        tipo = str(r.get("Tipo") or "").strip() or "Absoluta"
        longs, shorts = _split(r.get("Activos largo")), _split(r.get("Activos corto"))
        q = _num(r.get("Rendimiento esperado (%)"))
        if not longs and not shorts and np.isnan(q):
            continue  # fila vacía
        unknown = [a for a in longs + shorts if a not in tickers]
        if unknown:
            errors.append(f"{tag}: activos fuera del universo: {', '.join(unknown)}.")
            continue
        if set(longs) & set(shorts):
            errors.append(f"{tag}: un activo no puede estar en ambos lados.")
            continue
        if np.isnan(q):
            errors.append(f"{tag}: falta el rendimiento esperado.")
            continue
        if tipo == "Absoluta" and (len(longs) != 1 or shorts):
            errors.append(f"{tag}: una view absoluta requiere exactamente un activo en 'largo' y ninguno en 'corto'.")
            continue
        if tipo == "Relativa" and (len(longs) != 1 or len(shorts) != 1):
            errors.append(f"{tag}: una view relativa requiere un activo en 'largo' y uno en 'corto'.")
            continue
        if tipo == "Canasta" and not longs:
            errors.append(f"{tag}: una canasta requiere al menos un activo en 'largo'.")
            continue

        weighting = str(r.get("Ponderación canasta") or "Igual")

        def side_weights(assets):
            if weighting == "Capitalización" and all(caps.get(a) for a in assets):
                c = np.array([caps[a] for a in assets], float)
                return c / c.sum()
            if weighting == "Capitalización":
                warns.append(f"{tag}: sin capitalización para todos los activos; canasta a pesos iguales.")
            return np.full(len(assets), 1 / len(assets))

        p = pd.Series(0.0, index=tickers)
        p[longs] = side_weights(longs)
        if shorts:
            p[shorts] = -side_weights(shorts)
        is_absolute = not shorts
        q_dec = q / 100.0
        rows_p.append(p)
        q_in.append(q_dec)
        q_ex.append(q_dec - rf if is_absolute else q_dec)

        c = _num(r.get("Confianza (%)"), 50.0) / 100.0
        if not 0 < c <= 1:
            warns.append(f"{tag}: confianza {c * 100:.1f}% fuera de (0, 100]; se acota.")
        conf.append(float(np.clip(c, CONF_MIN, CONF_MAX)))
        l_, u_ = _num(r.get("Límite inferior (%)")) / 100, _num(r.get("Límite superior (%)")) / 100
        lo.append(l_)
        up.append(u_)
        if not np.isnan(l_) and not np.isnan(u_) and not (l_ <= q_dec <= u_):
            warns.append(f"{tag}: el rendimiento esperado no está dentro del intervalo [inf, sup].")
        desc = str(r.get("Descripción") or "").strip()
        labels.append(f"V{len(labels) + 1}: " + (desc if desc and desc != "nan" else
                      " + ".join(longs) + (" vs " + " + ".join(shorts) if shorts else "")))

    k_idx = labels
    P = pd.DataFrame(rows_p, index=k_idx, columns=tickers) if rows_p else pd.DataFrame(columns=tickers, dtype=float)
    mk = lambda v: pd.Series(v, index=k_idx, dtype=float)
    vs = ViewSet(P=P, Q=mk(q_ex), Q_input=mk(q_in), confidence=mk(conf), lower=mk(lo), upper=mk(up),
                 labels=labels, errors=errors, warnings=warns)
    if vs.k > 0:
        rank = np.linalg.matrix_rank(P.values)
        if rank < vs.k:
            vs.warnings.append(f"Las views son linealmente dependientes (rango {rank} < {vs.k}); "
                               "el modelo sigue siendo válido con Ω > 0, pero revise la redundancia.")
    return vs


# -----------------------------------------------------------------------------
# Ω
# -----------------------------------------------------------------------------
def omega_he_litterman(P: np.ndarray, cov: np.ndarray, tau: float) -> np.ndarray:
    return np.diag(np.diag(P @ (tau * cov) @ P.T))


def omega_idzorek(P: np.ndarray, cov: np.ndarray, tau: float, confidence: np.ndarray) -> np.ndarray:
    c = np.clip(np.asarray(confidence, float), CONF_MIN, CONF_MAX)
    alpha = (1 - c) / c
    return np.diag(alpha * np.diag(P @ (tau * cov) @ P.T))


def omega_interval(lower: np.ndarray, upper: np.ndarray, prob: float) -> np.ndarray:
    lower, upper = np.asarray(lower, float), np.asarray(upper, float)
    if np.any(np.isnan(lower)) or np.any(np.isnan(upper)):
        raise ValueError("El método de intervalos requiere límite inferior y superior en todas las views activas.")
    if np.any(upper <= lower):
        raise ValueError("En cada view el límite superior debe ser mayor que el inferior.")
    z = stats.norm.ppf(0.5 + prob / 2)
    return np.diag(((upper - lower) / (2 * z)) ** 2)


def idzorek_numeric(p: np.ndarray, q: float, cov: np.ndarray, pi: np.ndarray, w_mkt: np.ndarray,
                    tau: float, delta: float, confidence: float) -> float:
    """Calibración original de Idzorek (2005) para una view: min ‖w(ω) − w_objetivo‖²."""
    p = np.atleast_2d(p)
    ts = tau * cov
    inv_dc = np.linalg.inv(delta * cov)

    def weights(omega):
        a = p @ ts @ p.T + omega
        mu = pi + (ts @ p.T @ np.linalg.solve(a, np.atleast_1d(q - p @ pi))).ravel()
        return inv_dc @ mu

    w100 = weights(np.zeros((1, 1)))
    target = w_mkt + confidence * (w100 - w_mkt)
    base = float((p @ ts @ p.T).item())
    obj = lambda lw: float(np.sum((weights(np.array([[np.exp(lw)]])) - target) ** 2))
    res = optimize.minimize_scalar(obj, bounds=(np.log(base * 1e-8), np.log(base * 1e8)),
                                   method="bounded", options={"xatol": 1e-12})
    return float(np.exp(res.x))


def build_omega(method: str, vs: ViewSet, cov: pd.DataFrame, tau: float,
                interval_prob: float = 0.9) -> np.ndarray:
    if vs.k == 0:
        return np.zeros((0, 0))
    P, S = vs.P.values, cov.values
    if method == "he_litterman":
        return omega_he_litterman(P, S, tau)
    if method == "interval":
        return omega_interval(vs.lower.values, vs.upper.values, interval_prob)
    return omega_idzorek(P, S, tau, vs.confidence.values)


def implied_confidence(P: np.ndarray, cov: np.ndarray, tau: float, omega: np.ndarray) -> np.ndarray:
    """Confianza equivalente de Idzorek para cualquier Ω: c = pτΣpᵀ / (pτΣpᵀ + ω)."""
    base = np.diag(P @ (tau * cov) @ P.T)
    return base / (base + np.diag(omega))


# ==============================================================================
# DISTRIBUCIÓN POSTERIOR BLACK-LITTERMAN
# ==============================================================================
# Distribución posterior de Black-Litterman (forma de He & Litterman, 1999).
#
#     A      = P τΣ Pᵀ + Ω
#     μ_BL   = Π + τΣ Pᵀ A⁻¹ (Q − P Π)
#     M      = τΣ − τΣ Pᵀ A⁻¹ P τΣ              (incertidumbre de la media posterior)
#     Σ_BL   = Σ + M
#
# Esta forma no requiere invertir Ω, por lo que admite confianza del 100 % (Ω = 0).
# Sin views: μ_BL = Π y Σ_BL = (1 + τ) Σ.
@dataclass
class PosteriorResult:
    mu_bl: pd.Series                 # exceso anual
    cov_bl: pd.DataFrame
    M: pd.DataFrame
    view_contrib: pd.DataFrame       # activos × views: aporte de cada view a μ_BL − Π
    view_gap: pd.Series              # Q − PΠ (sorpresa de cada view respecto del equilibrio)
    view_posterior: pd.Series        # P μ_BL
    omega: np.ndarray


def black_litterman(pi: pd.Series, cov: pd.DataFrame, tau: float, P: pd.DataFrame | None = None,
                    Q: pd.Series | None = None, omega: np.ndarray | None = None) -> PosteriorResult:
    assets = cov.index
    S = cov.values
    ts = tau * S
    pi_v = pi.reindex(assets).values
    if P is None or len(P) == 0:
        M = ts.copy()
        return PosteriorResult(mu_bl=pi.reindex(assets).copy(),
                               cov_bl=pd.DataFrame(S + M, index=assets, columns=assets),
                               M=pd.DataFrame(M, index=assets, columns=assets),
                               view_contrib=pd.DataFrame(index=assets), view_gap=pd.Series(dtype=float),
                               view_posterior=pd.Series(dtype=float), omega=np.zeros((0, 0)))
    Pm = P.reindex(columns=assets).values
    q = Q.values
    A = Pm @ ts @ Pm.T + omega
    A = (A + A.T) / 2
    gap = q - Pm @ pi_v
    K = ts @ Pm.T @ np.linalg.pinv(A)           # n × k  (ganancia de Kalman)
    contrib = K * gap[None, :]                  # descomposición aditiva exacta
    mu = pi_v + contrib.sum(axis=1)
    M = ts - K @ Pm @ ts
    M = (M + M.T) / 2
    cov_bl = S + M
    return PosteriorResult(mu_bl=pd.Series(mu, index=assets),
                           cov_bl=pd.DataFrame(cov_bl, index=assets, columns=assets),
                           M=pd.DataFrame(M, index=assets, columns=assets),
                           view_contrib=pd.DataFrame(contrib, index=assets, columns=P.index),
                           view_gap=pd.Series(gap, index=P.index),
                           view_posterior=pd.Series(Pm @ mu, index=P.index), omega=omega)


def unconstrained_weights(mu: pd.Series, cov: pd.DataFrame, delta: float) -> pd.Series:
    """w* = (δΣ)⁻¹ μ  — solución analítica sin restricciones."""
    return pd.Series(np.linalg.solve(delta * cov.values, mu.reindex(cov.index).values), index=cov.index)


def weight_contributions(contrib: pd.DataFrame, cov: pd.DataFrame, delta: float) -> pd.DataFrame:
    """Aporte de cada view a los pesos sin restricciones (lineal en μ para Σ dado)."""
    if contrib.empty:
        return contrib
    inv = np.linalg.inv(delta * cov.values)
    return pd.DataFrame(inv @ contrib.values, index=contrib.index, columns=contrib.columns)


# ==============================================================================
# OPTIMIZACIÓN MEDIA-VARIANZA (cvxpy)
# ==============================================================================
# Optimización media-varianza (utilidad cuadrática) con restricciones convexas:
#
#     max_w   μᵀw − (δ/2) wᵀ Σ w
#     s.a.    presupuesto (Σw = 1 | Σw ≤ 1 | libre), w ≥ 0 (opcional),
#             l_i ≤ w_i ≤ u_i, L_g ≤ Σ_{i∈g} w_i ≤ U_g, ‖w − w_0‖₁ ≤ T
#
# μ se expresa en exceso de r_f: el efectivo (Σw < 1) rinde r_f, por lo que la
# formulación es correcta para los tres modos de presupuesto.
# Resuelto con cvxpy (CLARABEL/OSQP/SCS).
@dataclass
class Constraints:
    budget: str = "full"                        # full | no_leverage | free
    long_only: bool = True
    lower: pd.Series | None = None              # decimal
    upper: pd.Series | None = None
    groups: dict[str, str] = field(default_factory=dict)            # ticker → grupo
    group_bounds: dict[str, tuple[float, float]] = field(default_factory=dict)  # decimal
    turnover_max: float | None = None           # decimal (‖Δw‖₁)
    current: pd.Series | None = None


@dataclass
class OptResult:
    weights: pd.Series
    status: str
    objective: float
    ok: bool
    message: str = ""


def _psd(S: np.ndarray):
    return cp.psd_wrap((S + S.T) / 2)


def _constraint_list(w, assets, c: Constraints):
    cons = []
    if c.budget == "full":
        cons.append(cp.sum(w) == 1)
    elif c.budget == "no_leverage":
        cons.append(cp.sum(w) <= 1)
    if c.long_only:
        cons.append(w >= 0)
    if c.lower is not None:
        lo = c.lower.reindex(assets).values.astype(float)
        idx = [i for i in range(len(assets)) if np.isfinite(lo[i])]
        if idx:
            cons.append(w[idx] >= lo[idx])
    if c.upper is not None:
        hi = c.upper.reindex(assets).values.astype(float)
        idx = [i for i in range(len(assets)) if np.isfinite(hi[i])]
        if idx:
            cons.append(w[idx] <= hi[idx])
    for g, (lo, hi) in c.group_bounds.items():
        idx = [i for i, a in enumerate(assets) if c.groups.get(a) == g]
        if not idx:
            continue
        s = cp.sum(w[idx])
        if lo is not None and np.isfinite(lo):
            cons.append(s >= lo)
        if hi is not None and np.isfinite(hi):
            cons.append(s <= hi)
    if c.turnover_max is not None and c.current is not None:
        cons.append(cp.norm1(w - c.current.reindex(assets).fillna(0).values) <= c.turnover_max)
    return cons


def _solve(prob: cp.Problem) -> str:
    for solver in ("CLARABEL", "OSQP", "SCS"):
        if solver not in cp.installed_solvers():
            continue
        try:
            prob.solve(solver=solver)
            if prob.status in ("optimal", "optimal_inaccurate"):
                return prob.status
        except cp.SolverError:
            continue
    return prob.status or "solver_error"


def mean_variance(mu: pd.Series, cov: pd.DataFrame, delta: float, c: Constraints) -> OptResult:
    assets = list(cov.index)
    n = len(assets)
    w = cp.Variable(n)
    obj = cp.Maximize(mu.reindex(assets).values @ w - (delta / 2) * cp.quad_form(w, _psd(cov.values)))
    prob = cp.Problem(obj, _constraint_list(w, assets, c))
    status = _solve(prob)
    if status not in ("optimal", "optimal_inaccurate") or w.value is None:
        return OptResult(pd.Series(np.nan, index=assets), status, float("nan"), False,
                         "Problema no factible o no acotado: revise que las restricciones sean compatibles "
                         "(p. ej., suma de mínimos ≤ 100% ≤ suma de máximos).")
    wv = np.where(np.abs(w.value) < 1e-9, 0.0, w.value)
    return OptResult(pd.Series(wv, index=assets), status, float(prob.value), True)


def _min_var_for_target(mu, cov, c, target):
    assets = list(cov.index)
    w = cp.Variable(len(assets))
    cons = _constraint_list(w, assets, c) + [mu.reindex(assets).values @ w >= target]
    prob = cp.Problem(cp.Minimize(cp.quad_form(w, _psd(cov.values))), cons)
    _solve(prob)
    return None if w.value is None or prob.status not in ("optimal", "optimal_inaccurate") else w.value


def efficient_frontier(mu: pd.Series, cov: pd.DataFrame, c: Constraints, n_points: int = 40) -> pd.DataFrame:
    """Frontera eficiente (Σw = 1) bajo las mismas restricciones de pesos. μ en exceso."""
    fc = Constraints(budget="full", long_only=c.long_only, lower=c.lower, upper=c.upper,
                     groups=c.groups, group_bounds=c.group_bounds)
    assets = list(cov.index)
    m = mu.reindex(assets).values
    S = cov.values
    # extremos: mínima varianza y máximo rendimiento factible
    w = cp.Variable(len(assets))
    cons = _constraint_list(w, assets, fc)
    p_min = cp.Problem(cp.Minimize(cp.quad_form(w, _psd(S))), cons)
    _solve(p_min)
    if w.value is None:
        return pd.DataFrame(columns=["ret", "vol"])
    r_min = float(m @ w.value)
    w2 = cp.Variable(len(assets))
    p_max = cp.Problem(cp.Maximize(m @ w2), _constraint_list(w2, assets, fc))
    _solve(p_max)
    if w2.value is None:  # rendimiento no acotado (sin long-only ni límites)
        r_max = r_min + 3 * np.sqrt(np.max(np.diag(S)))
    else:
        r_max = float(m @ w2.value)
    rows = []
    for t in np.linspace(r_min, r_max, n_points):
        wv = _min_var_for_target(mu, cov, fc, t)
        if wv is not None:
            rows.append({"ret": float(m @ wv), "vol": float(np.sqrt(max(wv @ S @ wv, 0))),
                         **{a: float(x) for a, x in zip(assets, wv)}})
    return pd.DataFrame(rows).drop_duplicates(subset=["ret"]).sort_values("vol").reset_index(drop=True)


def portfolio_point(w: pd.Series, mu: pd.Series, cov: pd.DataFrame) -> tuple[float, float]:
    wv = w.reindex(cov.index).fillna(0).values
    return float(mu.reindex(cov.index).values @ wv), float(np.sqrt(wv @ cov.values @ wv))


# ==============================================================================
# ORQUESTADOR DEL MODELO
# ==============================================================================
# Orquestador: ejecuta el modelo Black-Litterman completo a partir de datos y configuración.
@dataclass
class ModelResult:
    cfg: ModelConfig
    rf: float
    cov: CovResult
    hist: pd.DataFrame
    w_mkt: pd.Series
    wmkt_method_used: str
    delta: float
    delta_raw: float | None
    tau: float
    tau_note: str
    pi: pd.Series
    views: ViewSet
    omega: np.ndarray
    implied_conf: pd.Series
    post: PosteriorResult
    sigma_opt: pd.DataFrame
    constraints: Constraints
    opt: OptResult
    w_unconstrained: pd.Series
    w_unconstrained_contrib: pd.DataFrame
    consistency_gap: float
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def w_opt(self) -> pd.Series:
        return self.opt.weights


def resolve_rf(cfg: ModelConfig, md: MarketData) -> float:
    if cfg.rf_method == "manual":
        return cfg.rf_manual_pct / 100
    if cfg.rf_method == "irx_mean":
        return md.rf_annual_mean
    return md.rf_annual_last


def build_constraints(cfg: ModelConfig, assets: list[str], w_ref: pd.Series) -> Constraints:
    lower = upper = None
    if cfg.use_asset_bounds:
        lower = pd.Series({a: cfg.asset_bounds.get(a, [cfg.default_min_pct, cfg.default_max_pct])[0] / 100
                           for a in assets})
        upper = pd.Series({a: cfg.asset_bounds.get(a, [cfg.default_min_pct, cfg.default_max_pct])[1] / 100
                           for a in assets})
    groups, gb = {}, {}
    if cfg.use_groups:
        groups = {a: cfg.asset_groups.get(a, "") for a in assets if cfg.asset_groups.get(a)}
        gb = {g: (b[0] / 100 if b[0] is not None else None, b[1] / 100 if b[1] is not None else None)
              for g, b in cfg.group_bounds.items()}
    current = None
    if cfg.use_turnover:
        if cfg.current_weights and sum(abs(v) for v in cfg.current_weights.values()) > 0:
            current = pd.Series({a: cfg.current_weights.get(a, 0.0) / 100 for a in assets})
        else:
            current = w_ref.reindex(assets)
    return Constraints(budget=cfg.budget, long_only=cfg.long_only, lower=lower, upper=upper,
                       groups=groups, group_bounds=gb,
                       turnover_max=cfg.turnover_max_pct / 100 if cfg.use_turnover else None,
                       current=current)


def run_model(md: MarketData, cfg: ModelConfig, views_df: pd.DataFrame | None,
              returns: pd.DataFrame | None = None, bench_returns: pd.Series | None = None,
              rf_daily: pd.Series | None = None, rf_override: float | None = None) -> ModelResult:
    """Ejecuta el pipeline completo. Los argumentos opcionales permiten ventanas (backtest)."""
    warns: list[str] = []
    R = md.returns if returns is None else returns
    bR = md.benchmark_returns if bench_returns is None else bench_returns
    rfd = md.rf_daily if rf_daily is None else rf_daily
    assets = list(R.columns)
    rf = resolve_rf(cfg, md) if rf_override is None else rf_override

    cov = estimate_covariance(R, cfg.cov_method, cfg.ewma_lambda)
    if cov.psd_repaired:
        warns.append("Σ no era semidefinida positiva; se reparó recortando eigenvalores.")

    w_mkt, used, w_w = market_weights(assets, cfg.wmkt_method, md.market_caps_usd, cfg.manual_weights)
    warns += w_w
    delta, delta_raw, d_w = resolve_delta(cfg.delta_method, cfg.delta_manual, bR, rfd, cfg.delta_bounds)
    warns += d_w
    tau, tau_note = resolve_tau(cfg.tau_method, cfg.tau_manual, len(R) / 252)
    pi = implied_returns(cov.cov, w_mkt, delta)

    vs = build_views(views_df, assets, rf, md.market_caps_usd)
    errors = list(vs.errors)
    warns += vs.warnings
    try:
        omega = build_omega(cfg.omega_method, vs, cov.cov, tau, cfg.interval_prob_pct / 100)
    except ValueError as e:
        errors.append(str(e))
        omega = build_omega("he_litterman", vs, cov.cov, tau)
        warns.append("Se usó Ω de He-Litterman como respaldo por el error anterior.")
    post = black_litterman(pi, cov.cov, tau, vs.P if vs.k else None, vs.Q if vs.k else None, omega)
    ic = (pd.Series(implied_confidence(vs.P.values, cov.cov.values, tau, omega), index=vs.labels)
          if vs.k else pd.Series(dtype=float))

    sigma_opt = post.cov_bl if cfg.sigma_opt == "posterior" else cov.cov
    cons = build_constraints(cfg, assets, w_mkt)
    opt = mean_variance(post.mu_bl, sigma_opt, delta, cons)
    if not opt.ok:
        errors.append(opt.message)
    w_unc = unconstrained_weights(post.mu_bl, sigma_opt, delta)
    w_unc_contrib = weight_contributions(post.view_contrib, sigma_opt, delta)

    # Control de consistencia: sin views y sin restricciones, w* ∝ w_mkt
    # (con Σ_BL sin views = (1+τ)Σ ⇒ w* = w_mkt/(1+τ); con Σ ⇒ w* = w_mkt)
    sigma_nv = cov.cov * (1 + tau) if cfg.sigma_opt == "posterior" else cov.cov
    w0 = unconstrained_weights(pi, sigma_nv, delta)
    consistency = float(np.max(np.abs(w0 / w0.sum() - w_mkt))) if abs(w0.sum()) > 1e-12 else float("nan")

    return ModelResult(cfg=cfg, rf=rf, cov=cov, hist=annualized_stats(R), w_mkt=w_mkt,
                       wmkt_method_used=used, delta=delta, delta_raw=delta_raw, tau=tau,
                       tau_note=tau_note, pi=pi, views=vs, omega=omega, implied_conf=ic, post=post,
                       sigma_opt=sigma_opt, constraints=cons, opt=opt, w_unconstrained=w_unc,
                       w_unconstrained_contrib=w_unc_contrib, consistency_gap=consistency,
                       warnings=warns, errors=errors)


# ==============================================================================
# MÉTRICAS DE RIESGO
# ==============================================================================
# Métricas de riesgo.
#
# Ex-ante (con μ_BL y la Σ usada en la optimización):
#     rendimiento esperado, volatilidad, Sharpe, VaR/CVaR paramétricos (normal),
#     contribución al riesgo  RC_i = w_i (Σw)_i / σ_p   (Σ_i RC_i = σ_p, Euler).
# Ex-post (pesos fijos sobre la historia; solo descriptivo, sujeto a sesgo in-sample):
#     VaR/CVaR históricos y Cornish-Fisher, beta, tracking error, drawdown.
def portfolio_returns(returns: pd.DataFrame, w: pd.Series, rf_daily: pd.Series | None = None) -> pd.Series:
    wv = w.reindex(returns.columns).fillna(0)
    r = returns @ wv
    cash = 1 - wv.sum()
    if rf_daily is not None and abs(cash) > 1e-12:
        r = r + cash * rf_daily.reindex(r.index).fillna(0)
    return r


def risk_contributions(w: pd.Series, cov: pd.DataFrame) -> pd.DataFrame:
    wv = w.reindex(cov.index).fillna(0).values
    S = cov.values
    sigma = float(np.sqrt(wv @ S @ wv))
    mrc = S @ wv / sigma if sigma > 0 else np.zeros_like(wv)
    rc = wv * mrc
    return pd.DataFrame({"Peso": wv, "Riesgo marginal": mrc, "Contribución (vol)": rc,
                         "Contribución (%)": rc / sigma if sigma > 0 else rc}, index=cov.index)


def parametric_var(mu_ann: float, vol_ann: float, alpha: float = 0.95, horizon_days: int = 1):
    h = horizon_days / TRADING_DAYS
    m, s = mu_ann * h, vol_ann * np.sqrt(h)
    z = stats.norm.ppf(alpha)
    var = -(m - z * s)
    cvar = -(m - s * stats.norm.pdf(z) / (1 - alpha))
    return float(var), float(cvar)


def historical_var(r: pd.Series, alpha: float = 0.95):
    q = np.quantile(r.dropna(), 1 - alpha)
    return float(-q), float(-r[r <= q].mean())


def cornish_fisher_var(r: pd.Series, alpha: float = 0.95) -> float:
    z = stats.norm.ppf(1 - alpha)
    s, k = stats.skew(r.dropna()), stats.kurtosis(r.dropna())
    zcf = z + (z**2 - 1) * s / 6 + (z**3 - 3 * z) * k / 24 - (2 * z**3 - 5 * z) * s**2 / 36
    return float(-(r.mean() + zcf * r.std(ddof=1)))


def drawdown(r: pd.Series) -> pd.Series:
    v = (1 + r).cumprod()
    return v / v.cummax() - 1


def performance_stats(r: pd.Series, rf_daily: pd.Series | None = None, bench: pd.Series | None = None) -> dict:
    r = r.dropna()
    n = len(r)
    if n < 2:
        return {}
    rf = rf_daily.reindex(r.index).fillna(0) if rf_daily is not None else pd.Series(0.0, index=r.index)
    ex = r - rf
    years = n / TRADING_DAYS
    total = float((1 + r).prod())
    cagr = total ** (1 / years) - 1
    vol = r.std(ddof=1) * np.sqrt(TRADING_DAYS)
    downside = np.sqrt((np.minimum(ex, 0) ** 2).mean()) * np.sqrt(TRADING_DAYS)
    dd = drawdown(r)
    out = {"Rendimiento total": total - 1, "CAGR": cagr, "Volatilidad anual": vol,
           "Sharpe": ex.mean() * TRADING_DAYS / vol if vol > 0 else np.nan,
           "Sortino": ex.mean() * TRADING_DAYS / downside if downside > 0 else np.nan,
           "Máx. drawdown": dd.min(), "Calmar": cagr / abs(dd.min()) if dd.min() < 0 else np.nan}
    if bench is not None:
        b = bench.reindex(r.index)
        a = (r - b).dropna()
        te = a.std(ddof=1) * np.sqrt(TRADING_DAYS)
        out.update({"Beta": float(np.cov(r.loc[a.index], b.loc[a.index])[0, 1] / b.loc[a.index].var(ddof=1)),
                    "Tracking error": te,
                    "Information ratio": a.mean() * TRADING_DAYS / te if te > 0 else np.nan})
    return out


def ex_ante_summary(w: pd.Series, mu_excess: pd.Series, cov: pd.DataFrame, rf: float,
                    alpha: float = 0.95) -> dict:
    wv = w.reindex(cov.index).fillna(0)
    ex = float(mu_excess.reindex(cov.index).values @ wv.values)
    vol = float(np.sqrt(wv.values @ cov.values @ wv.values))
    total = rf + ex
    var1, cvar1 = parametric_var(total, vol, alpha, 1)
    var_y, cvar_y = parametric_var(total, vol, alpha, TRADING_DAYS)
    sd = np.sqrt(np.diag(cov.values))
    return {"Rendimiento esperado": total, "Prima sobre r_f": ex, "Volatilidad": vol,
            "Sharpe": ex / vol if vol > 0 else np.nan,
            f"VaR {alpha:.0%} 1 día": var1, f"CVaR {alpha:.0%} 1 día": cvar1,
            f"VaR {alpha:.0%} 1 año": var_y, f"CVaR {alpha:.0%} 1 año": cvar_y,
            "Ratio de diversificación": float(np.abs(wv.values) @ sd / vol) if vol > 0 else np.nan,
            "N efectivo (1/Σw²)": float(1 / (wv**2).sum()) if (wv**2).sum() > 0 else np.nan,
            "Inversión total": float(wv.sum())}


def ex_post_summary(returns: pd.DataFrame, w: pd.Series, rf_daily: pd.Series, bench: pd.Series,
                    alpha: float = 0.95) -> dict:
    r = portfolio_returns(returns, w, rf_daily)
    hv, hcv = historical_var(r, alpha)
    out = performance_stats(r, rf_daily, bench)
    out.update({f"VaR histórico {alpha:.0%} 1 día": hv, f"CVaR histórico {alpha:.0%} 1 día": hcv,
                f"VaR Cornish-Fisher {alpha:.0%} 1 día": cornish_fisher_var(r, alpha),
                "Asimetría": float(stats.skew(r)), "Curtosis en exceso": float(stats.kurtosis(r))})
    return out


# ==============================================================================
# BACKTEST WALK-FORWARD
# ==============================================================================
# Backtesting walk-forward (fuera de muestra para Σ, δ, Π y r_f).
#
# En cada fecha de rebalanceo t se estiman Σ, δ (si es implícito) y Π usando
# únicamente la ventana [t − L, t]; se optimiza con las mismas restricciones y la
# cartera se mantiene (con deriva de precios) hasta el siguiente rebalanceo.
# Los costos de transacción se cargan como  costo = bps × ‖w_nuevo − w_derivado‖₁.
#
# Sesgos que se declaran explícitamente
# -------------------------------------
# * Views: son opiniones actuales; aplicarlas fijas en el pasado introduce sesgo
#   de anticipación (look-ahead). Por eso se reporta también la estrategia de
#   equilibrio (sin views), que no lo tiene.
# * Capitalización de mercado: Yahoo solo ofrece la actual; si w_mkt usa
#   capitalización, los pesos de referencia históricos también tienen look-ahead.
@dataclass
class BacktestResult:
    returns: pd.DataFrame                  # diarios por estrategia
    weights: dict[str, pd.DataFrame]       # pesos objetivo en cada rebalanceo
    turnover: pd.DataFrame
    stats: pd.DataFrame
    notes: list[str] = field(default_factory=list)


def rebalance_dates(index: pd.DatetimeIndex, freq: str, lookback: int) -> list[pd.Timestamp]:
    s = pd.Series(index, index=index)
    key = index.to_period("M" if freq == "M" else "Q")
    last = s.groupby(key).max().tolist()
    return [d for d in last if index.get_loc(d) >= lookback - 1 and d != index[-1]]


def _simulate(returns: pd.DataFrame, rf_daily: pd.Series, targets: dict[pd.Timestamp, pd.Series],
              cost_bps: float) -> tuple[pd.Series, pd.Series]:
    dates = sorted(targets)
    start = returns.index.get_loc(dates[0]) + 1
    idx = returns.index[start:]
    assets = list(returns.columns)
    w = targets[dates[0]].reindex(assets).fillna(0).values.astype(float)
    cash = 1 - w.sum()
    out, tos = [], {dates[0]: float(np.abs(w).sum() + abs(cash))}
    pending_cost = cost_bps / 1e4 * tos[dates[0]]
    R, RF = returns.values, rf_daily.reindex(returns.index).fillna(0).values
    for i in range(start, len(returns)):
        r = R[i]
        rp = float(w @ r + cash * RF[i]) - pending_cost
        pending_cost = 0.0
        gross = 1 + float(w @ r + cash * RF[i])
        w = w * (1 + r) / gross
        cash = cash * (1 + RF[i]) / gross
        d = returns.index[i]
        if d in targets:
            new = targets[d].reindex(assets).fillna(0).values.astype(float)
            to = float(np.abs(new - w).sum() + abs((1 - new.sum()) - cash))
            tos[d] = to
            pending_cost = cost_bps / 1e4 * to
            w, cash = new, 1 - new.sum()
        out.append(rp)
    return pd.Series(out, index=idx), pd.Series(tos)


def run_backtest(md: MarketData, cfg: ModelConfig, views_df: pd.DataFrame | None, freq: str = "M",
                 lookback: int = TRADING_DAYS, cost_bps: float = 10.0, progress=None) -> BacktestResult:
    R, bR, rfd = md.returns, md.benchmark_returns, md.rf_daily
    dates = rebalance_dates(R.index, freq, lookback)
    if len(dates) < 2:
        raise ValueError("Datos insuficientes para el backtest: amplíe el periodo o reduzca la ventana de estimación.")
    notes = []
    has_views = views_df is not None and len(views_df) > 0
    if has_views:
        notes.append("La estrategia 'BL con views' aplica las views actuales en todo el historial: "
                     "incluye sesgo de anticipación (look-ahead) y no debe leerse como desempeño alcanzable.")
    if cfg.wmkt_method == "market_cap":
        notes.append("w_mkt usa la capitalización actual (Yahoo no ofrece historial de capitalización): "
                     "el portafolio de referencia histórico también tiene look-ahead.")

    targets = {"BL con views": {}, "Equilibrio (sin views)": {}, "1/N": {}}
    prev = {k: None for k in targets}
    n = len(R.columns)
    for j, d in enumerate(dates):
        pos = R.index.get_loc(d)
        win = R.iloc[pos - lookback + 1: pos + 1]
        rf_at = float(rfd.iloc[pos] * TRADING_DAYS) if cfg.rf_method != "manual" else cfg.rf_manual_pct / 100
        for name, vdf in (("BL con views", views_df), ("Equilibrio (sin views)", None)):
            if name == "BL con views" and not has_views:
                continue
            c = cfg
            if cfg.use_turnover and prev[name] is not None:
                c = replace(cfg, current_weights={a: float(v) * 100 for a, v in prev[name].items()})
            res = run_model(md, c, vdf, returns=win, bench_returns=bR.loc[win.index],
                            rf_daily=rfd.loc[win.index], rf_override=rf_at)
            w = res.w_opt if res.opt.ok else (prev[name] if prev[name] is not None else res.w_mkt)
            targets[name][d] = w
            prev[name] = w
        targets["1/N"][d] = pd.Series(1 / n, index=R.columns)
        if progress:
            progress((j + 1) / len(dates))

    rets, tos, wts = {}, {}, {}
    for name, tg in targets.items():
        if not tg:
            continue
        rets[name], tos[name] = _simulate(R, rfd, tg, cost_bps)
        wts[name] = pd.DataFrame(tg).T
    rets_df = pd.DataFrame(rets)
    rets_df["Índice de referencia"] = bR.reindex(rets_df.index)
    stats = {k: performance_stats(rets_df[k], rfd, rets_df.iloc[:, -1]) for k in rets_df.columns}
    st = pd.DataFrame(stats)
    to_df = pd.DataFrame(tos)
    st.loc["Rotación promedio por rebalanceo"] = [to_df[c].iloc[1:].mean() if c in to_df else np.nan
                                                  for c in st.columns]
    return BacktestResult(returns=rets_df, weights=wts, turnover=to_df, stats=st, notes=notes)


# ==============================================================================
# EXPORTACIÓN (EXCEL, JSON, PDF)
# ==============================================================================
# Exportación de resultados: Excel (tablas y parámetros), JSON (configuración) y PDF (reporte).
PCT_HINT = ("Rend", "Prima", "Vol", "VaR", "CVaR", "Peso", "Π", "μ", "CAGR", "drawdown",
            "Tracking", "Contribución (%)", "Q", "Rotación", "Inversión", "w_", "Mín", "Máx")


PCT_SHEETS = {"Rendimientos", "Pesos", "Views", "Aporte views a μ", "Aporte views a w",
              "Contribución riesgo", "Frontera"}
NON_PCT_COLS = {"Divisa original", "Cap. mercado USD (mm)", "ω (varianza de la view)"}

# =============================================================================
# Tablas comunes
# =============================================================================
def results_tables(res, md, risk_ante: dict, risk_post: dict, frontier: pd.DataFrame | None = None,
                   backtest=None) -> dict[str, pd.DataFrame]:
    rf = res.rf
    assets = list(res.cov.cov.index)
    returns_tbl = pd.DataFrame({
        "Divisa original": pd.Series(md.currencies),
        "Cap. mercado USD (mm)": pd.Series({a: (md.market_caps_usd.get(a) or np.nan) / 1e6 for a in assets}),
        "w_mkt": res.w_mkt,
        "Rend. histórico (aritm.)": res.hist["Rend. aritmético anual"],
        "Π equilibrio (total)": res.pi + rf,
        "μ_BL posterior (total)": res.post.mu_bl + rf,
        "μ_BL − Π": res.post.mu_bl - res.pi,
        "Vol. histórica": res.hist["Volatilidad anual"],
        "Vol. posterior": np.sqrt(np.diag(res.post.cov_bl.values)),
    }).reindex(assets)
    weights_tbl = pd.DataFrame({"w_mkt (equilibrio)": res.w_mkt,
                                "w* sin restricciones": res.w_unconstrained,
                                "w* óptimo (restringido)": res.w_opt,
                                "Activo vs mercado": res.w_opt - res.w_mkt}).reindex(assets)
    views_tbl = pd.DataFrame()
    if res.views.k:
        views_tbl = pd.DataFrame({"Q capturado": res.views.Q_input, "Q exceso r_f": res.views.Q,
                                  "PΠ (equilibrio)": res.views.P.values @ res.pi.values,
                                  "Sorpresa Q − PΠ": res.post.view_gap,
                                  "Pμ_BL (posterior)": res.post.view_posterior,
                                  "Confianza capturada": res.views.confidence,
                                  "Confianza implícita": res.implied_conf,
                                  "ω (varianza de la view)": np.diag(res.omega)})
    params = pd.DataFrame({"Valor": {
        "Fecha de ejecución": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Tickers": ", ".join(assets), "Periodo": f"{md.returns.index[0].date()} a {md.returns.index[-1].date()}",
        "Observaciones diarias": md.n_obs, "Índice de referencia": res.cfg.benchmark,
        "Moneda base": "USD", "Tasa libre de riesgo (anual)": rf,
        "Pesos de equilibrio": res.wmkt_method_used, "δ (aversión al riesgo)": res.delta,
        "δ implícito sin acotar": res.delta_raw, "τ": res.tau, "Criterio τ": res.tau_note,
        "Covarianza": res.cov.method, "Intensidad de shrinkage": res.cov.shrinkage,
        "Número de condición Σ": res.cov.condition_number, "Método Ω": res.cfg.omega_method,
        "Σ para optimizar": res.cfg.sigma_opt, "Presupuesto": res.cfg.budget,
        "Solo largos": res.cfg.long_only, "Estado del optimizador": res.opt.status}})
    tables = {"Parámetros": params, "Rendimientos": returns_tbl, "Pesos": weights_tbl}
    if not views_tbl.empty:
        tables["Views"] = views_tbl
        tables["Matriz P"] = res.views.P
        tables["Aporte views a μ"] = res.post.view_contrib
        tables["Aporte views a w"] = res.w_unconstrained_contrib
    tables["Riesgo ex-ante"] = pd.DataFrame({"Valor": risk_ante})
    tables["Riesgo ex-post"] = pd.DataFrame({"Valor": risk_post})
    tables["Contribución riesgo"] = risk_contributions(res.w_opt, res.sigma_opt)
    tables["Σ histórica"] = res.cov.cov
    tables["Σ posterior"] = res.post.cov_bl
    tables["Correlaciones"] = cov_to_corr(res.cov.cov)
    if frontier is not None and not frontier.empty:
        f = frontier.copy()
        f["ret"] = f["ret"] + rf
        tables["Frontera"] = f.rename(columns={"ret": "Rendimiento", "vol": "Volatilidad"})
    if backtest is not None:
        tables["Backtest métricas"] = backtest.stats
        tables["Backtest valor"] = (1 + backtest.returns.fillna(0)).cumprod()
    tables["Precios USD"] = md.prices_usd
    return tables


# =============================================================================
# Excel
# =============================================================================
def to_excel(tables: dict[str, pd.DataFrame]) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in tables.items():
            sheet = name[:31]
            out = df.copy()
            if isinstance(out.index, pd.DatetimeIndex):
                out.index = out.index.date
            out.to_excel(xw, sheet_name=sheet)
            ws = xw.sheets[sheet]
            head = PatternFill("solid", fgColor="1F3A5F")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "B2"
            ws.column_dimensions["A"].width = 30
            for j, col in enumerate(out.columns, start=2):
                ws.column_dimensions[get_column_letter(j)].width = 16
                is_pct = name in PCT_SHEETS and str(col) not in NON_PCT_COLS
                for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    for c in row:
                        if isinstance(c.value, (int, float)):
                            c.number_format = "0.00%" if is_pct else "#,##0.0000"
            if name in ("Riesgo ex-ante", "Riesgo ex-post", "Parámetros", "Backtest métricas"):
                for row in ws.iter_rows(min_row=2):
                    label = str(row[0].value)
                    pct = any(h in label for h in PCT_HINT) and not any(
                        k in label for k in ("Sharpe", "Sortino", "Calmar", "Beta", "ratio", "N efectivo",
                                             "Número", "Observ", "δ", "τ", "Asimetría", "Curtosis"))
                    for c in row[1:]:
                        if isinstance(c.value, (int, float)):
                            c.number_format = "0.00%" if pct else "0.0000"
    return buf.getvalue()


# =============================================================================
# JSON
# =============================================================================
def config_to_json(cfg: ModelConfig) -> str:
    d = cfg.to_dict()
    d["_schema"] = "black-litterman-config/1.0"
    return json.dumps(d, ensure_ascii=False, indent=2, default=float)


def config_from_json(text: str) -> ModelConfig:
    d = json.loads(text)
    d.pop("_schema", None)
    return ModelConfig.from_dict(d)


def views_to_records(df: pd.DataFrame) -> list[dict]:
    clean = df.copy().replace({np.nan: None})
    return clean.to_dict(orient="records")


# =============================================================================
# PDF
# =============================================================================
METHODOLOGY = [
    ("1. Datos", "Precios de cierre ajustados de Yahoo Finance convertidos a USD con el tipo de cambio "
     "de Yahoo ({CCY}=X). Calendarios alineados (relleno hacia adelante de hasta 5 días por feriados "
     "locales) y rendimientos simples diarios anualizados con 252 días. Tasa libre de riesgo: ^IRX."),
    ("2. Covarianza", "Σ anualizada estimada con el método elegido: muestral, Ledoit-Wolf (2004) con "
     "objetivo de correlación constante e intensidad óptima, o EWMA RiskMetrics (λ)."),
    ("3. Equilibrio", "Rendimientos implícitos en exceso de r_f por optimización inversa: Π = δ Σ w_mkt."),
    ("4. Views", "P μ = Q + ε, ε ~ N(0, Ω). Views absolutas (Q − r_f), relativas y de canasta. Ω por "
     "He-Litterman (diag(PτΣPᵀ)), Idzorek (ω = (1−c)/c · pτΣpᵀ) o intervalos ((U−L)/2z)²."),
    ("5. Posterior", "μ_BL = Π + τΣPᵀ(PτΣPᵀ+Ω)⁻¹(Q − PΠ);  Σ_BL = Σ + τΣ − τΣPᵀ(PτΣPᵀ+Ω)⁻¹PτΣ."),
    ("6. Optimización", "max μ_BLᵀw − (δ/2) wᵀΣw sujeto a presupuesto, no negatividad, límites por "
     "activo y grupo, y rotación máxima; resuelto con cvxpy."),
]


def _register_fonts() -> tuple[str, str]:
    """Registra DejaVu Sans (incluida con matplotlib) para soportar letras griegas y superíndices."""
    from pathlib import Path
    import matplotlib
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    base = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    try:
        if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DejaVu", str(base / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(base / "DejaVuSans-Bold.ttf")))
            from reportlab.lib.fonts import addMapping
            addMapping("DejaVu", 0, 0, "DejaVu"); addMapping("DejaVu", 1, 0, "DejaVu-Bold")
            addMapping("DejaVu", 0, 1, "DejaVu"); addMapping("DejaVu", 1, 1, "DejaVu-Bold")
        return "DejaVu", "DejaVu-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


def _fig_to_img(fig, width_cm=16):
    from reportlab.lib.units import cm
    from reportlab.platypus import Image
    b = io.BytesIO()
    fig.savefig(b, format="png", dpi=160, bbox_inches="tight")
    b.seek(0)
    w, h = fig.get_size_inches()
    return Image(b, width=width_cm * cm, height=width_cm * cm * h / w)


def _charts(res, frontier, backtest):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    figs = []
    assets = list(res.cov.cov.index)
    x = np.arange(len(assets))
    fig, ax = plt.subplots(figsize=(8, 3.4))
    ax.bar(x - 0.2, res.w_mkt.values * 100, 0.4, label="w_mkt", color="#9AA5B1")
    ax.bar(x + 0.2, res.w_opt.reindex(assets).values * 100, 0.4, label="w* óptimo", color="#1F5FAD")
    ax.set_xticks(x, assets, rotation=30, ha="right", fontsize=8)
    ax.set_ylabel("%"); ax.set_title("Pesos: equilibrio vs óptimo", fontsize=10)
    ax.legend(frameon=False, fontsize=8); ax.spines[["top", "right"]].set_visible(False)
    figs.append(fig)
    if frontier is not None and not frontier.empty:
        fig, ax = plt.subplots(figsize=(8, 3.8))
        ax.plot(frontier["vol"] * 100, (frontier["ret"] + res.rf) * 100, color="#1F5FAD", lw=2,
                label="Frontera posterior")
        for w, lab, col in ((res.w_opt, "Óptimo", "#C0392B"), (res.w_mkt, "Mercado", "#555")):
            r, v = portfolio_point(w, res.post.mu_bl, res.sigma_opt)
            ax.scatter(v * 100, (r + res.rf) * 100, color=col, zorder=3, label=lab)
        ax.set_xlabel("Volatilidad (%)"); ax.set_ylabel("Rendimiento esperado (%)")
        ax.set_title("Frontera eficiente posterior", fontsize=10)
        ax.legend(frameon=False, fontsize=8); ax.spines[["top", "right"]].set_visible(False)
        figs.append(fig)
    if backtest is not None:
        fig, ax = plt.subplots(figsize=(8, 3.6))
        v = (1 + backtest.returns.fillna(0)).cumprod() * 100
        for c in v.columns:
            ax.plot(v.index, v[c], lw=1.4, label=c)
        ax.set_title("Backtest walk-forward (valor base 100)", fontsize=10)
        ax.legend(frameon=False, fontsize=7); ax.spines[["top", "right"]].set_visible(False)
        figs.append(fig)
    return figs


def _df_table(df: pd.DataFrame, pct_cols=None, fmt="{:.4f}", max_rows=30, font=("Helvetica", "Helvetica-Bold")):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    pct_cols = pct_cols or []
    data = [[""] + [str(c) for c in df.columns]]
    for idx, row in df.head(max_rows).iterrows():
        cells = [str(idx)]
        for c, v in row.items():
            if isinstance(v, (int, float, np.floating)) and not pd.isna(v):
                cells.append(f"{v:.2%}" if c in pct_cols or pct_cols == "all" else fmt.format(v))
            else:
                cells.append("" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v))
        data.append(cells)
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7), ("FONTNAME", (0, 0), (-1, -1), font[0]),
        ("FONTNAME", (0, 0), (-1, 0), font[1]),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F7")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9CED6"))]))
    return t


def to_pdf(res, md, risk_ante: dict, frontier: pd.DataFrame | None = None, backtest=None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title="Reporte Black-Litterman", author="Black-Litterman Optimizer")
    font = _register_fonts()
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontName=font[1])
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName=font[1])
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=8.5, leading=11, fontName=font[0])
    small = ParagraphStyle("s", parent=body, fontSize=7.5, textColor="#555555")
    rf = res.rf
    story = [Paragraph("Reporte de optimización Black-Litterman", h1),
             Paragraph(f"Generado el {datetime.now():%Y-%m-%d %H:%M} · Moneda base USD · "
                       f"Periodo {md.returns.index[0].date()} a {md.returns.index[-1].date()} "
                       f"({md.n_obs} observaciones diarias)", small), Spacer(1, 8)]

    story.append(Paragraph("Parámetros del modelo", h2))
    p = pd.DataFrame({"Valor": {
        "Universo": ", ".join(res.cov.cov.index), "Índice de referencia": res.cfg.benchmark,
        "Tasa libre de riesgo": f"{rf:.2%}", "Pesos de equilibrio": res.wmkt_method_used,
        "δ": f"{res.delta:.3f}", "τ": f"{res.tau:.4f} ({res.tau_note})",
        "Covarianza": res.cov.method + (f" (shrinkage {res.cov.shrinkage:.3f})" if res.cov.shrinkage is not None else ""),
        "Ω": res.cfg.omega_method, "Σ en optimización": res.cfg.sigma_opt,
        "Presupuesto / solo largos": f"{res.cfg.budget} / {res.cfg.long_only}"}})
    story += [_df_table(p, font=font), Spacer(1, 8)]

    story.append(Paragraph("Rendimientos esperados (anuales, totales)", h2))
    t = pd.DataFrame({"w_mkt": res.w_mkt, "Histórico": res.hist["Rend. aritmético anual"],
                      "Π + r_f": res.pi + rf, "μ_BL + r_f": res.post.mu_bl + rf,
                      "w* óptimo": res.w_opt})
    story += [_df_table(t, pct_cols="all", font=font), Spacer(1, 6)]
    if res.views.k:
        story.append(Paragraph("Views", h2))
        v = pd.DataFrame({"Q": res.views.Q_input, "Confianza": res.views.confidence,
                          "Conf. implícita": res.implied_conf, "Pμ_BL": res.post.view_posterior})
        story += [_df_table(v, pct_cols="all", font=font), Spacer(1, 6)]
    story.append(Paragraph("Riesgo ex-ante del portafolio óptimo", h2))
    ra = pd.DataFrame({"Valor": {k: (f"{v:.4f}" if any(s in k for s in ("Sharpe", "Ratio", "N efectivo"))
                                     else f"{v:.2%}") for k, v in risk_ante.items()}})
    story += [_df_table(ra, font=font), PageBreak()]

    story.append(Paragraph("Gráficas", h2))
    for fig in _charts(res, frontier, backtest):
        story += [_fig_to_img(fig), Spacer(1, 6)]
    if backtest is not None:
        story.append(Paragraph("Backtest walk-forward", h2))
        bs = backtest.stats.copy()
        story.append(_df_table(bs, fmt="{:.4f}", font=font))
        for n in backtest.notes:
            story.append(Paragraph("• " + n, small))
    story += [PageBreak(), Paragraph("Metodología", h2)]
    for title, text in METHODOLOGY:
        story += [Paragraph(f"<b>{title}.</b> {text}", body), Spacer(1, 4)]
    story.append(Paragraph("Referencias: Black & Litterman (1992); He & Litterman (1999); Idzorek (2005); "
                           "Ledoit & Wolf (2004); Walters (2014); J.P. Morgan RiskMetrics (1996).", small))
    story.append(Paragraph("Este reporte es una herramienta analítica y no constituye asesoría de inversión.", small))
    doc.build(story)
    return buf.getvalue()


# ==============================================================================
# GRÁFICAS (PLOTLY)
# ==============================================================================
# Gráficas Plotly para la app (paleta categórica fija, validada para daltonismo).
# 8 tonos validados + 2 de extensión para universos de 9-10 activos (máximo de la app)
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948",
          "#7a5230", "#0f7c8c"]
NEUTRAL = "#8a8984"
DIVERGING = [[0.0, "#e34948"], [0.5, "#f0efec"], [1.0, "#2a78d6"]]


def _layout(fig: go.Figure, title: str, ytitle: str = "", xtitle: str = "", pct_y=False, height=420):
    fig.update_layout(title=dict(text=title, x=0, y=0.98, yanchor="top", font=dict(size=15)), height=height,
                      margin=dict(l=10, r=10, t=85, b=10), hovermode="closest",
                      legend=dict(orientation="h", yanchor="bottom", y=1.01, xanchor="left", x=0),
                      bargap=0.25, bargroupgap=0.06)
    fig.update_xaxes(title=xtitle, showgrid=False, zeroline=False)
    fig.update_yaxes(title=ytitle, gridcolor="rgba(128,128,128,0.18)", zeroline=True,
                     zerolinecolor="rgba(128,128,128,0.45)", tickformat=".1%" if pct_y else None)
    return fig


def grouped_bars(df: pd.DataFrame, title: str, pct=True, ytitle="") -> go.Figure:
    fig = go.Figure()
    for i, c in enumerate(df.columns):
        fig.add_bar(x=df.index, y=df[c], name=str(c), marker_color=SERIES[i % len(SERIES)],
                    marker_line_width=0,
                    hovertemplate="%{x}<br>" + str(c) + ": %{y:" + (".2%" if pct else ".4f") + "}<extra></extra>")
    return _layout(fig, title, ytitle, pct_y=pct)


def stacked_contrib(df: pd.DataFrame, title: str, pct=True) -> go.Figure:
    """df: activos × views (aportes aditivos)."""
    fig = go.Figure()
    for i, c in enumerate(df.columns):
        fig.add_bar(x=df.index, y=df[c], name=str(c), marker_color=SERIES[i % len(SERIES)],
                    hovertemplate="%{x}<br>" + str(c) + ": %{y:" + (".2%" if pct else ".4f") + "}<extra></extra>")
    fig.update_layout(barmode="relative")
    return _layout(fig, title, pct_y=pct)


def price_index(prices: pd.DataFrame, bench: pd.Series | None = None, title="Precios en USD (base 100)"):
    base = prices / prices.iloc[0] * 100
    fig = go.Figure()
    for i, c in enumerate(base.columns):
        fig.add_scatter(x=base.index, y=base[c], name=c, mode="lines", line=dict(width=1.6, color=SERIES[i % len(SERIES)]),
                        hovertemplate=c + ": %{y:.1f}<extra></extra>")
    if bench is not None:
        b = bench / bench.iloc[0] * 100
        fig.add_scatter(x=b.index, y=b, name=str(bench.name or "Índice"), mode="lines",
                        line=dict(width=2, color=NEUTRAL, dash="dot"),
                        hovertemplate="Índice: %{y:.1f}<extra></extra>")
    fig.update_layout(hovermode="x unified")
    return _layout(fig, title, "Índice (base 100)")


def value_curves(returns: pd.DataFrame, title="Valor del portafolio (base 100)"):
    v = (1 + returns.fillna(0)).cumprod() * 100
    fig = go.Figure()
    for i, c in enumerate(v.columns):
        is_bench = "Índice" in c
        fig.add_scatter(x=v.index, y=v[c], name=c, mode="lines",
                        line=dict(width=2, color=NEUTRAL if is_bench else SERIES[i % len(SERIES)], dash="dot" if is_bench else None),
                        hovertemplate=c + ": %{y:.1f}<extra></extra>")
    fig.update_layout(hovermode="x unified")
    return _layout(fig, title, "Valor")


def drawdown_curves(returns: pd.DataFrame, title="Drawdown"):
    fig = go.Figure()
    for i, c in enumerate(returns.columns):
        v = (1 + returns[c].fillna(0)).cumprod()
        dd = v / v.cummax() - 1
        is_bench = "Índice" in c
        fig.add_scatter(x=dd.index, y=dd, name=c, mode="lines",
                        line=dict(width=1.5, color=NEUTRAL if is_bench else SERIES[i % len(SERIES)], dash="dot" if is_bench else None),
                        hovertemplate=c + ": %{y:.1%}<extra></extra>")
    fig.update_layout(hovermode="x unified")
    return _layout(fig, title, pct_y=True, height=320)


def heatmap(corr: pd.DataFrame, title="Matriz de correlaciones"):
    fig = go.Figure(go.Heatmap(z=corr.values, x=corr.columns, y=corr.index, zmin=-1, zmax=1,
                               colorscale=DIVERGING, text=np.round(corr.values, 2), texttemplate="%{text}",
                               hovertemplate="%{y} / %{x}: %{z:.3f}<extra></extra>", xgap=2, ygap=2))
    fig.update_yaxes(autorange="reversed")
    return _layout(fig, title, height=460)


def frontiers(curves: dict[str, pd.DataFrame], points: dict[str, tuple[float, float]],
              assets: dict[str, tuple[float, float]] | None = None, title="Fronteras eficientes"):
    fig = go.Figure()
    for i, (name, f) in enumerate(curves.items()):
        if f is None or f.empty:
            continue
        fig.add_scatter(x=f["vol"], y=f["ret"], name=name, mode="lines",
                        line=dict(width=2.2 if i == 0 else 1.6, color=SERIES[i], dash=None if i == 0 else "dash"),
                        hovertemplate=name + "<br>σ %{x:.2%} · E[R] %{y:.2%}<extra></extra>")
    symbols = ["star", "diamond", "square", "circle", "triangle-up"]
    for j, (name, (r, v)) in enumerate(points.items()):
        fig.add_scatter(x=[v], y=[r], name=name, mode="markers",
                        marker=dict(size=13 if j == 0 else 10, symbol=symbols[j % 5], color=SERIES[(j + 3) % len(SERIES)],
                                    line=dict(width=2, color="white")),
                        hovertemplate=name + "<br>σ %{x:.2%} · E[R] %{y:.2%}<extra></extra>")
    if assets:
        fig.add_scatter(x=[v for r, v in assets.values()], y=[r for r, v in assets.values()],
                        text=list(assets.keys()), mode="markers+text", textposition="middle right",
                        name="Activos (posterior)", marker=dict(size=8, color=NEUTRAL),
                        hovertemplate="%{text}<br>σ %{x:.2%} · E[R] %{y:.2%}<extra></extra>")
    fig = _layout(fig, title, "Rendimiento esperado anual", "Volatilidad anual", pct_y=True, height=520)
    fig.update_xaxes(tickformat=".0%")
    return fig


def weights_area(w: pd.DataFrame, title="Pesos objetivo en cada rebalanceo"):
    fig = go.Figure()
    for i, c in enumerate(w.columns):
        fig.add_scatter(x=w.index, y=w[c], name=c, stackgroup="one", mode="lines",
                        line=dict(width=0.5, color=SERIES[i % len(SERIES)]),
                        hovertemplate=c + ": %{y:.1%}<extra></extra>")
    fig.update_layout(hovermode="x unified")
    return _layout(fig, title, pct_y=True, height=380)


# ==============================================================================
# APLICACIÓN STREAMLIT
# ==============================================================================
st.set_page_config(page_title="Black-Litterman Optimizer", page_icon="📈", layout="wide")
try:  # acento institucional en controles nativos (radios, sliders, toggles, selecciones)
    st._config.set_option("theme.primaryColor", "#1F5FAD")
except Exception:  # noqa: BLE001
    pass

INV = lambda d: {v: k for k, v in d.items()}  # clave interna → etiqueta
PCT = "{:.2%}"

# =============================================================================
# Catálogo de activos, universos sugeridos e índices
# =============================================================================
CATALOG = {
    # EE. UU.
    "AAPL": ("Apple", "Tecnología"), "MSFT": ("Microsoft", "Tecnología"), "NVDA": ("NVIDIA", "Tecnología"),
    "GOOGL": ("Alphabet", "Comunicación"), "AMZN": ("Amazon", "Consumo discrecional"),
    "META": ("Meta Platforms", "Comunicación"), "TSLA": ("Tesla", "Consumo discrecional"),
    "AVGO": ("Broadcom", "Tecnología"), "JPM": ("JPMorgan Chase", "Financiero"),
    "BAC": ("Bank of America", "Financiero"), "V": ("Visa", "Financiero"),
    "BRK-B": ("Berkshire Hathaway", "Financiero"), "JNJ": ("Johnson & Johnson", "Salud"),
    "UNH": ("UnitedHealth", "Salud"), "LLY": ("Eli Lilly", "Salud"), "PG": ("Procter & Gamble", "Consumo básico"),
    "KO": ("Coca-Cola", "Consumo básico"), "WMT": ("Walmart", "Consumo básico"), "XOM": ("Exxon Mobil", "Energía"),
    "CVX": ("Chevron", "Energía"), "CAT": ("Caterpillar", "Industrial"), "NEE": ("NextEra Energy", "Servicios públicos"),
    # México (BMV)
    "WALMEX.MX": ("Walmart de México", "Consumo básico"), "GMEXICOB.MX": ("Grupo México", "Materiales"),
    "AMXB.MX": ("América Móvil", "Comunicación"), "FEMSAUBD.MX": ("FEMSA", "Consumo básico"),
    "GFNORTEO.MX": ("Banorte", "Financiero"), "CEMEXCPO.MX": ("Cemex", "Materiales"),
    "BIMBOA.MX": ("Grupo Bimbo", "Consumo básico"), "GAPB.MX": ("GAP Aeropuertos", "Industrial"),
    "ASURB.MX": ("ASUR Aeropuertos", "Industrial"), "KOFUBL.MX": ("Coca-Cola FEMSA", "Consumo básico"),
    "ORBIA.MX": ("Orbia", "Materiales"), "PINFRA.MX": ("Pinfra", "Industrial"),
    # ETFs
    "SPY": ("ETF S&P 500", "Acciones EE. UU."), "QQQ": ("ETF Nasdaq 100", "Acciones EE. UU."),
    "IWM": ("ETF Russell 2000", "Acciones EE. UU."), "EFA": ("ETF Mercados desarrollados", "Acciones internacionales"),
    "EEM": ("ETF Mercados emergentes", "Acciones internacionales"), "EWW": ("ETF MSCI México", "Acciones internacionales"),
    "AGG": ("ETF Bonos agregados EE. UU.", "Renta fija"), "TLT": ("ETF Tesoros 20+ años", "Renta fija"),
    "LQD": ("ETF Bonos corporativos", "Renta fija"), "GLD": ("ETF Oro", "Materias primas"),
    "VNQ": ("ETF Bienes raíces", "Bienes raíces"),
}
GROUPS = sorted({g for _, g in CATALOG.values()} | {"Otro"})
PRESETS = {
    "EE. UU. + México": ["AAPL", "MSFT", "NVDA", "JPM", "XOM", "JNJ", "WALMEX.MX", "GMEXICOB.MX"],
    "Tecnología EE. UU.": ["AAPL", "MSFT", "NVDA", "GOOGL", "AMZN", "META", "AVGO"],
    "México (BMV)": ["WALMEX.MX", "GMEXICOB.MX", "AMXB.MX", "FEMSAUBD.MX", "GFNORTEO.MX", "CEMEXCPO.MX",
                     "GAPB.MX", "BIMBOA.MX"],
    "Multiactivo (ETFs)": ["SPY", "EFA", "EEM", "AGG", "TLT", "GLD", "VNQ"],
    "Defensivo": ["JNJ", "PG", "KO", "WMT", "NEE", "BRK-B", "AGG"],
}
BENCHMARKS = {"^GSPC": "S&P 500", "^NDX": "Nasdaq 100", "^DJI": "Dow Jones Industrial", "ACWI": "MSCI ACWI (ETF)",
              "URTH": "MSCI World (ETF)", "^MXX": "S&P/BMV IPC (México)", "EEM": "MSCI Emergentes (ETF)"}
PERIODS = {"1 año": 1, "3 años": 3, "5 años": 5, "10 años": 10, "Personalizado": None}
SHORT_LABELS = {  # etiquetas compactas para los radios de la barra lateral
    "Capitalización de mercado": "Capitalización", "Pesos iguales (1/N)": "Iguales (1/N)", "Pesos manuales": "Manuales",
    "Fijo estándar (2.5)": "Estándar 2.5", "Implícito del mercado": "Implícito", "Manual": "Manual",
    "Ledoit-Wolf (correlación constante)": "Ledoit-Wolf", "Muestral": "Muestral", "EWMA (RiskMetrics)": "EWMA",
    "Estándar (0.05)": "0.05", "1/T (años de datos)": "1/T",
    "Idzorek (confianza %)": "Idzorek (confianza %)", "He-Litterman (proporcional a la varianza)": "He-Litterman",
    "Intervalos de confianza": "Intervalos",
    "Σ_BL posterior (He-Litterman)": "Σ posterior (BL)", "Σ histórica (prior)": "Σ histórica",
    "Totalmente invertido (Σw = 1)": "100 % invertido", "Sin apalancamiento (Σw ≤ 1, resto en T-Bills)": "Permite efectivo",
    "Libre (permite apalancamiento)": "Libre (apalancamiento)",
    "^IRX último dato (T-Bill 13 semanas)": "T-Bill 13s (último)", "^IRX promedio del periodo": "T-Bill 13s (promedio)",
}
short = lambda x: SHORT_LABELS.get(x, x)
ticker_label = lambda t: f"{t} · {CATALOG[t][0]}" if t in CATALOG else t

# =============================================================================
# Estilo
# =============================================================================
ACCENT = "#1F5FAD"
CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
.stApp { font-family: 'Inter', system-ui, -apple-system, sans-serif; }
.block-container { padding-top: 1.4rem; padding-bottom: 3rem; max-width: 1440px; }
header[data-testid="stHeader"] { background: transparent; }

/* Encabezado */
.bl-hero { background: linear-gradient(120deg, #0B1F3A 0%, #143A6E 55%, #1F5FAD 100%);
  border-radius: 16px; padding: 26px 30px 22px; margin-bottom: 18px; color: #fff;
  box-shadow: 0 8px 24px rgba(11,31,58,.18); }
.bl-hero .bl-title { font-size: 1.75rem; font-weight: 700; letter-spacing: -.015em; line-height: 1.2; }
.bl-hero .bl-sub { color: rgba(255,255,255,.80); font-size: .95rem; margin: 4px 0 14px; }
.bl-chips { display: flex; flex-wrap: wrap; gap: 8px; }
.bl-chip { background: rgba(255,255,255,.12); border: 1px solid rgba(255,255,255,.22); border-radius: 999px;
  padding: 4px 12px; font-size: .78rem; color: #fff; white-space: nowrap; }
.bl-chip b { font-weight: 600; }

/* Tarjetas de indicadores */
.bl-kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(138px, 1fr)); gap: 12px; margin: 4px 0 14px; }
.bl-kpi { border: 1px solid rgba(128,128,128,.22); border-left: 4px solid #1F5FAD; border-radius: 12px;
  padding: 12px 16px; background: rgba(31,95,173,.045); }
.bl-kpi.pos { border-left-color: #1A8F5A; } .bl-kpi.neg { border-left-color: #C0392B; }
.bl-kpi.neu { border-left-color: #8A8984; }
.bl-kpi-label { font-size: .78rem; opacity: .70; font-weight: 600; white-space: nowrap; overflow: hidden;
  text-overflow: ellipsis; }
.bl-kpi-value { font-size: 1.35rem; font-weight: 700; margin-top: 2px; font-variant-numeric: tabular-nums; }
.bl-kpi-sub { font-size: .76rem; opacity: .62; margin-top: 1px; }

/* Títulos de sección */
.bl-section { margin: 18px 0 8px; padding-bottom: 6px; border-bottom: 1px solid rgba(128,128,128,.22); }
.bl-section .t { font-size: 1.08rem; font-weight: 650; }
.bl-section .d { font-size: .84rem; opacity: .68; margin-top: 2px; }
.bl-note { font-size: .82rem; opacity: .72; }
.bl-step { display: inline-block; background: #1F5FAD; color: #fff; border-radius: 6px; font-size: .72rem;
  font-weight: 700; padding: 1px 7px; margin-right: 6px; }

/* Pestañas */
.stTabs [data-baseweb="tab-list"] { gap: 4px; border-bottom: 1px solid rgba(128,128,128,.25); }
.stTabs [data-baseweb="tab"] { height: 2.7rem; padding: 0 1.05rem; border-radius: 10px 10px 0 0; font-weight: 600; }
.stTabs [aria-selected="true"] { background: rgba(31,95,173,.10); }
.stTabs [aria-selected="true"] p { color: #1F5FAD; }
.stTabs [data-baseweb="tab-highlight"] { background-color: #1F5FAD; }

/* Botones */
.stButton > button, .stDownloadButton > button { border-radius: 9px; font-weight: 600; }
.stButton > button[kind="primary"], .stDownloadButton > button[kind="primary"] {
  background: #1F5FAD; border-color: #1F5FAD; color: #fff; }
.stButton > button[kind="primary"]:hover, .stDownloadButton > button[kind="primary"]:hover {
  background: #174C8C; border-color: #174C8C; color: #fff; }

/* Acentos institucionales (azul) en controles */
span[data-baseweb="tag"] { background-color: #1F5FAD !important; }
[data-testid="stSlider"] [role="slider"] { background-color: #1F5FAD !important; }
[data-testid="stSliderThumbValue"] { color: #1F5FAD !important; }

/* Barra lateral */
section[data-testid="stSidebar"] { background: linear-gradient(180deg, #0B1F3A 0%, #10294D 100%); }
section[data-testid="stSidebar"] [data-testid="stWidgetLabel"] p,
section[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p,
section[data-testid="stSidebar"] [data-testid="stRadio"] label p,
section[data-testid="stSidebar"] [data-testid="stCheckbox"] label p,
section[data-testid="stSidebar"] summary p, section[data-testid="stSidebar"] summary span,
section[data-testid="stSidebar"] h1, section[data-testid="stSidebar"] h2, section[data-testid="stSidebar"] h3 {
  color: #E9EFF8 !important; }
section[data-testid="stSidebar"] [data-testid="stCaptionContainer"] p { color: rgba(233,239,248,.62) !important; }
section[data-testid="stSidebar"] [data-testid="stExpander"] details {
  border: 1px solid rgba(255,255,255,.14); border-radius: 12px; background: rgba(255,255,255,.03); }
section[data-testid="stSidebar"] [data-testid="stSliderTickBarMin"],
section[data-testid="stSidebar"] [data-testid="stSliderTickBarMax"] { color: rgba(233,239,248,.6); }
section[data-testid="stSidebar"] hr { border-color: rgba(255,255,255,.15); }
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button {
  background: rgba(255,255,255,.08); border: 1px solid rgba(255,255,255,.22); }
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button p { color: #E9EFF8 !important; }
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button:hover { border-color: #9FC1EE; }
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button[aria-checked="true"],
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button[data-selected="true"] {
  background: #1F5FAD; border-color: #5B8FD6; }
section[data-testid="stSidebar"] [data-testid="stButtonGroup"] button[aria-checked="true"] p { color: #fff !important; }
section[data-testid="stSidebar"] [data-testid="stExpander"] summary { background: transparent !important; }
section[data-testid="stSidebar"] [data-testid="stExpander"] summary:hover p { color: #9FC1EE !important; }
section[data-testid="stSidebar"] [data-testid="stTooltipIcon"] svg { color: rgba(233,239,248,.55); stroke: rgba(233,239,248,.55); }
.bl-brand { color: #fff; font-weight: 700; font-size: 1.15rem; letter-spacing: -.01em; margin: -8px 0 2px; }
.bl-brand-sub { color: rgba(233,239,248,.6); font-size: .78rem; margin-bottom: 10px; }
.bl-side-h { color: #9FC1EE; font-size: .72rem; font-weight: 700; text-transform: uppercase; letter-spacing: .08em;
  margin: 14px 0 4px; }
.bl-pending { background: rgba(237,161,0,.16); border: 1px solid rgba(237,161,0,.45); color: #F6D48A;
  border-radius: 8px; padding: 6px 10px; font-size: .78rem; margin: 6px 0; }

/* Tablas */
[data-testid="stDataFrame"], [data-testid="stDataEditor"] { border-radius: 10px; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


def section(title: str, desc: str | None = None, step: str | None = None):
    badge = f'<span class="bl-step">{step}</span>' if step else ""
    st.markdown(f'<div class="bl-section"><div class="t">{badge}{title}</div>'
                + (f'<div class="d">{desc}</div>' if desc else "") + "</div>", unsafe_allow_html=True)


def kpis(items: list[tuple], container=st):
    """items: (etiqueta, valor, subtítulo, tono) con tono en {'', 'pos', 'neg', 'neu'}."""
    cells = "".join(f'<div class="bl-kpi {tone}"><div class="bl-kpi-label">{lab}</div>'
                    f'<div class="bl-kpi-value">{val}</div><div class="bl-kpi-sub">{sub}</div></div>'
                    for lab, val, sub, tone in items)
    container.markdown(f'<div class="bl-kpis">{cells}</div>', unsafe_allow_html=True)


# =============================================================================
# Estado
# =============================================================================
def _period_dates(label: str) -> tuple[date, date]:
    end = date.today()
    years = PERIODS.get(label) or 5
    return end - timedelta(days=int(years * 365.25)), end


def _init_state():
    d = ModelConfig()
    start, end = _period_dates("5 años")
    defaults = {
        "k_tickers": PRESETS["EE. UU. + México"], "k_preset": None, "k_period": "5 años",
        "k_start": start, "k_end": end, "k_bench": d.benchmark, "k_missing": d.max_missing_pct,
        "k_rf_method": INV(RF_METHODS)[d.rf_method], "k_rf_manual": d.rf_manual_pct,
        "k_wmkt": INV(WMKT_METHODS)[d.wmkt_method],
        "k_delta": INV(DELTA_METHODS)[d.delta_method], "k_delta_manual": d.delta_manual,
        "k_delta_lo": d.delta_bounds[0], "k_delta_hi": d.delta_bounds[1],
        "k_cov": INV(COV_METHODS)[d.cov_method], "k_ewma": d.ewma_lambda,
        "k_tau": INV(TAU_METHODS)[d.tau_method], "k_tau_manual": d.tau_manual,
        "k_omega": INV(OMEGA_METHODS)[d.omega_method], "k_int_prob": d.interval_prob_pct,
        "k_sigma_opt": INV(SIGMA_OPT)[d.sigma_opt], "k_budget": INV(BUDGET_MODES)[d.budget],
        "k_long_only": d.long_only, "k_use_bounds": d.use_asset_bounds,
        "k_wrange": (d.default_min_pct, d.default_max_pct), "k_custom_bounds": False,
        "k_use_groups": d.use_groups, "k_use_turnover": d.use_turnover, "k_turnover": d.turnover_max_pct,
        "views_df": None, "assets_df": None, "groups_df": None,
        "ed_ver": 0, "data_params": None, "backtest": None, "exports": None, "auto_load": True,
    }
    for k, v in defaults.items():
        st.session_state.setdefault(k, v)


def _apply_preset():
    choice = st.session_state.k_preset
    if choice:
        st.session_state.k_tickers = list(PRESETS[choice])
        st.session_state.auto_load = True
    st.session_state.k_preset = None


def _on_period():
    if st.session_state.k_period and st.session_state.k_period != "Personalizado":
        st.session_state.k_start, st.session_state.k_end = _period_dates(st.session_state.k_period)
        st.session_state.auto_load = True


@st.cache_data(ttl=3600, show_spinner=False)
def cached_market_data(tickers: tuple, start: str, end: str, bench: str, missing: float):
    return load_market_data(list(tickers), start, end, bench, missing)


@st.cache_data(show_spinner=False, max_entries=64)
def cached_frontier(key: str, _mu, _cov, _cons):
    return efficient_frontier(_mu, _cov, _cons, n_points=35)


def _hash(*objs) -> str:
    return hashlib.sha1(pickle.dumps(objs)).hexdigest()


def _sync_assets_df(tickers: list[str]):
    s = st.session_state
    cols = ["Ticker", "Nombre", "Grupo", "Mín (%)", "Máx (%)", "Peso actual (%)", "Peso manual w_mkt (%)"]
    old = s.assets_df.set_index("Ticker") if s.assets_df is not None else pd.DataFrame(columns=cols[1:])
    if s.assets_df is not None and list(s.assets_df["Ticker"]) == tickers:
        return
    lo, hi = s.k_wrange
    rows = []
    for t in tickers:
        if t in old.index:
            rows.append({"Ticker": t, **old.loc[t].to_dict()})
        else:
            rows.append({"Ticker": t, "Nombre": CATALOG.get(t, (t, ""))[0], "Grupo": CATALOG.get(t, ("", "Otro"))[1],
                         "Mín (%)": float(lo), "Máx (%)": float(hi), "Peso actual (%)": 0.0,
                         "Peso manual w_mkt (%)": round(100 / len(tickers), 2)})
    s.assets_df = pd.DataFrame(rows, columns=cols)
    # conservar solo las views cuyos activos siguen en el universo; si no queda ninguna, usar ejemplos
    if s.views_df is not None:
        def _ok(row):
            names = [x.strip() for x in f"{row['Activos largo']},{row['Activos corto']}".split(",") if x.strip()
                     and x.strip() != "nan"]
            return all(n in tickers for n in names)
        kept = s.views_df[s.views_df.apply(_ok, axis=1)] if len(s.views_df) else s.views_df
        s.views_df = kept.reset_index(drop=True) if len(kept) else example_views_df(tickers)
    s.pop("views_current", None)
    s.ed_ver += 1


def tab_widget(fn, label, name, container=st, **kw):
    """Widget cuyo valor vive en session_state[name]; la clave cambia con ed_ver para forzar el valor mostrado."""
    val = getattr(container, fn)(label, value=st.session_state[name],
                                 key=f"w_{name}_{st.session_state.ed_ver}", **kw)
    st.session_state[name] = val
    return val


def fmt_pct_df(df: pd.DataFrame, pct_cols=None):
    cols = pct_cols if pct_cols is not None else df.select_dtypes("number").columns
    return df.style.format({c: PCT for c in cols}, na_rep="—")


_init_state()
s = st.session_state

# =============================================================================
# Barra lateral
# =============================================================================
with st.sidebar:
    st.markdown('<div class="bl-brand">📈 Black-Litterman</div>'
                '<div class="bl-brand-sub">Optimizador de portafolios · USD · Yahoo Finance</div>',
                unsafe_allow_html=True)

    st.markdown('<div class="bl-side-h">Universo de inversión</div>', unsafe_allow_html=True)
    st.pills("Universos sugeridos (un clic)", list(PRESETS), key="k_preset", on_change=_apply_preset)
    options = list(CATALOG) + [t for t in s.k_tickers if t not in CATALOG]
    st.multiselect("Activos (2 a 10)", options, key="k_tickers", format_func=ticker_label,
                   max_selections=MAX_ASSETS, accept_new_options=True,
                   placeholder="Busca por ticker o nombre…",
                   help="Elige del catálogo o escribe cualquier ticker de Yahoo Finance y pulsa Enter "
                        "(p. ej. SAP.DE, 7203.T, ORCL).")
    st.selectbox("Índice de referencia", list(BENCHMARKS), key="k_bench",
                 format_func=lambda t: f"{BENCHMARKS[t]} ({t})")
    st.segmented_control("Periodo histórico", list(PERIODS), key="k_period", on_change=_on_period,
                         required=True, width="stretch")
    if s.k_period == "Personalizado":
        c1, c2 = st.columns(2)
        tab_widget("date_input", "Inicio", "k_start", c1, max_value=date.today())
        tab_widget("date_input", "Fin", "k_end", c2, max_value=date.today())

    tickers_in = [t.strip().upper() for t in s.k_tickers if str(t).strip()]
    tickers_in = list(dict.fromkeys(tickers_in))
    params = (tuple(tickers_in), s.k_start.isoformat(), s.k_end.isoformat(), s.k_bench, float(s.k_missing))
    pending_box = st.empty()
    load_clicked = st.button("Actualizar datos", type="primary", width="stretch", icon=":material/refresh:")

    st.markdown('<div class="bl-side-h">Modelo</div>', unsafe_allow_html=True)
    with st.expander("Equilibrio de mercado", expanded=True, icon=":material/balance:"):
        st.radio("Pesos de referencia (w_mkt)", list(WMKT_METHODS), key="k_wmkt", format_func=short,
                 help="Portafolio de partida del que se infieren los rendimientos de equilibrio Π = δΣw_mkt.")
        st.radio("Aversión al riesgo (δ)", list(DELTA_METHODS), key="k_delta", format_func=short, horizontal=True,
                 help="Estándar: 2.5 (He-Litterman). Implícito: (E[Rm] − rf)/σ²m del índice, acotado.")
        if DELTA_METHODS[s.k_delta] == "manual":
            tab_widget("slider", "δ manual", "k_delta_manual", min_value=0.5, max_value=10.0, step=0.1)
        if DELTA_METHODS[s.k_delta] == "implied":
            c1, c2 = st.columns(2)
            tab_widget("number_input", "δ mínimo", "k_delta_lo", c1, min_value=0.1, max_value=5.0, step=0.1)
            tab_widget("number_input", "δ máximo", "k_delta_hi", c2, min_value=1.0, max_value=20.0, step=0.5)
    with st.expander("Riesgo e incertidumbre", expanded=True, icon=":material/tune:"):
        st.radio("Covarianza (Σ)", list(COV_METHODS), key="k_cov", format_func=short, horizontal=True,
                 help="Ledoit-Wolf reduce el ruido de estimación; EWMA pondera más lo reciente.")
        if COV_METHODS[s.k_cov] == "ewma":
            tab_widget("slider", "λ (decaimiento EWMA)", "k_ewma", min_value=0.80, max_value=0.995, step=0.005)
        st.radio("Escala de incertidumbre (τ)", list(TAU_METHODS), key="k_tau", format_func=short, horizontal=True,
                 help="0.05 es el estándar de la literatura. 1/T usa los años de datos.")
        if TAU_METHODS[s.k_tau] == "manual":
            tab_widget("slider", "τ manual", "k_tau_manual", min_value=0.01, max_value=1.0, step=0.01)
        st.radio("Confianza en views (Ω)", list(OMEGA_METHODS), key="k_omega", format_func=short,
                 help="Idzorek: confianza de 0 a 100 % por view. He-Litterman: proporcional a la varianza. "
                      "Intervalos: a partir de un rango ± alrededor del rendimiento esperado.")
        if OMEGA_METHODS[s.k_omega] == "interval":
            tab_widget("slider", "Probabilidad del intervalo (%)", "k_int_prob", min_value=50.0, max_value=99.0,
                       step=1.0)
    with st.expander("Optimización", expanded=False, icon=":material/target:"):
        st.caption("Objetivo: max μ_BLᵀw − (δ/2)·wᵀΣw")
        st.radio("Covarianza para optimizar", list(SIGMA_OPT), key="k_sigma_opt", format_func=short, horizontal=True)
        st.radio("Presupuesto", list(BUDGET_MODES), key="k_budget", format_func=short)
        st.toggle("Solo posiciones largas (w ≥ 0)", key="k_long_only")
    with st.expander("Datos avanzados", expanded=False, icon=":material/settings:"):
        st.radio("Tasa libre de riesgo (USD)", list(RF_METHODS), key="k_rf_method", format_func=short)
        if RF_METHODS[s.k_rf_method] == "manual":
            tab_widget("number_input", "r_f anual (%)", "k_rf_manual", min_value=0.0, max_value=20.0, step=0.05)
        st.slider("Máximo de datos faltantes por activo (%)", 0.0, 50.0, key="k_missing", step=1.0)
    st.caption(f"v{__version__} · Herramienta analítica; no constituye asesoría de inversión.")

# =============================================================================
# Carga de datos
# =============================================================================
if load_clicked or s.auto_load:
    if len(tickers_in) == 0:
        st.sidebar.error("Selecciona al menos un activo.")
    elif s.k_start >= s.k_end:
        st.sidebar.error("La fecha de inicio debe ser anterior a la fecha fin.")
    else:
        s.data_params = params
        s.backtest = None
        s.exports = None
    s.auto_load = False

if s.data_params is not None and params != s.data_params:
    pending_box.markdown('<div class="bl-pending">Hay cambios sin aplicar en el universo o el periodo. '
                         'Pulsa <b>Actualizar datos</b>.</div>', unsafe_allow_html=True)

if s.data_params is None:
    st.info("Selecciona tus activos en la barra lateral y pulsa **Actualizar datos**.")
    st.stop()

try:
    with st.spinner("Descargando datos de Yahoo Finance…"):
        md = cached_market_data(*s.data_params)
except Exception as e:  # noqa: BLE001
    st.error(f"No fue posible obtener los datos de Yahoo Finance: {e}")
    st.caption("Yahoo Finance puede limitar temporalmente las consultas. Espera unos segundos y pulsa "
               "**Actualizar datos** de nuevo, o revisa que los tickers existan en finance.yahoo.com.")
    st.stop()

assets = md.tickers
_sync_assets_df(assets)
if s.views_df is None:
    s.views_df = example_views_df(assets)
if s.groups_df is None:
    s.groups_df = pd.DataFrame({"Grupo": pd.Series(dtype=str), "Mín (%)": pd.Series(dtype=float),
                                "Máx (%)": pd.Series(dtype=float)})

bench_name = BENCHMARKS.get(s.data_params[3], s.data_params[3])
st.markdown(
    f'''<div class="bl-hero"><div class="bl-title">Optimizador de portafolios Black-Litterman</div>
    <div class="bl-sub">Equilibrio de mercado + opiniones del inversionista → rendimientos posteriores →
    portafolio óptimo con restricciones</div><div class="bl-chips">
    <span class="bl-chip"><b>{len(assets)}</b> activos</span>
    <span class="bl-chip">Periodo <b>{md.returns.index[0]:%d %b %Y} – {md.returns.index[-1]:%d %b %Y}</b></span>
    <span class="bl-chip">Índice <b>{bench_name}</b></span>
    <span class="bl-chip">Moneda base <b>USD</b></span>
    <span class="bl-chip">Fuente <b>Yahoo Finance</b></span></div></div>''', unsafe_allow_html=True)

kpi_box = st.container()
for t, why in md.dropped.items():
    st.warning(f"**{t}** se excluyó: {why}.")
for w in md.warnings:
    st.info(w)

tabs = st.tabs([":material/database: Datos", ":material/balance: Equilibrio", ":material/lightbulb: Views",
                ":material/target: Optimización", ":material/shield: Riesgo", ":material/history: Backtest",
                ":material/download: Exportar"])
t_data, t_eq, t_views, t_opt, t_risk, t_bt, t_exp = tabs

# ---------------------------------------------------------------------------- Datos
with t_data:
    section("Universo y calidad de datos", "Precios de cierre ajustados convertidos a USD.")
    info = pd.DataFrame({"Nombre": pd.Series({a: CATALOG.get(a, (md.names.get(a, a),))[0] for a in assets}),
                         "Divisa original": pd.Series(md.currencies),
                         "Cap. mercado USD (mm)": pd.Series({a: (md.market_caps_usd.get(a) or np.nan) / 1e6 for a in assets}),
                         "Datos faltantes (%)": pd.Series({a: md.missing_pct.get(a, 0) for a in assets})}).reindex(assets)
    l, r = st.columns([1.6, 1])
    l.dataframe(info.style.format({"Cap. mercado USD (mm)": "{:,.0f}", "Datos faltantes (%)": "{:.1f}"}, na_rep="—"),
                width="stretch")
    with r:
        kpis([("Observaciones", f"{md.n_obs:,}", "rendimientos diarios", "neu"),
              ("Años de datos", f"{md.years:.1f}", "ventana de estimación", "neu"),
              ("T-Bill 13 semanas", PCT.format(md.rf_annual_last), "último dato ^IRX", "neu")])
    st.plotly_chart(price_index(md.prices_usd, md.benchmark_prices.rename(bench_name)), width="stretch")
    section("Estadísticos históricos", "Anualizados, en USD.")
    hist_tbl = annualized_stats(md.returns)
    st.dataframe(fmt_pct_df(hist_tbl), width="stretch")

if len(assets) < MIN_ASSETS_OPT:
    for tab in tabs[1:]:
        with tab:
            st.info("Se requieren al menos 2 activos válidos para optimizar. Con un solo activo solo está disponible "
                    "el análisis individual de la pestaña **Datos**.")
    with t_data:
        st.plotly_chart(drawdown_curves(md.returns.iloc[:, 0].to_frame(assets[0])), width="stretch")
    st.stop()

# =============================================================================
# Entradas en pestañas (se leen antes de ejecutar el modelo)
# =============================================================================
ver = s.ed_ver

with t_eq:
    if WMKT_METHODS[s.k_wmkt] == "manual":
        section("Pesos manuales del portafolio de referencia", "Se normalizan automáticamente a 100 %.")
        man = st.data_editor(s.assets_df[["Ticker", "Nombre", "Peso manual w_mkt (%)"]], hide_index=True,
                             disabled=["Ticker", "Nombre"], key=f"ed_manual_{ver}", width="stretch",
                             column_config={"Peso manual w_mkt (%)": st.column_config.NumberColumn(
                                 min_value=0.0, max_value=100.0, step=0.5, format="%.2f")})
        s.assets_df["Peso manual w_mkt (%)"] = man["Peso manual w_mkt (%)"].values

with t_views:
    section("Construye tus views", "Elige el tipo de opinión, los activos y el rendimiento esperado; después "
            "pulsa Agregar. Las views se combinan con el equilibrio según su nivel de confianza.", "1")
    b_col, t_col = st.columns([1, 1.7], gap="large")
    with b_col:
        with st.container(border=True):
            vtype = st.segmented_control("Tipo de view", ["Absoluta", "Relativa", "Canasta"], default="Absoluta",
                                         key="vb_tipo", required=True, width="stretch")
            longs, shorts, weighting = [], [], "Igual"
            if vtype == "Absoluta":
                a = st.selectbox("Activo", assets, format_func=ticker_label, key="vb_abs")
                longs = [a]
                q_label, q_default = "Rendimiento total anual esperado (%)", 10.0
            elif vtype == "Relativa":
                c1, c2 = st.columns(2)
                a = c1.selectbox("Activo que supera", assets, format_func=ticker_label, key="vb_rel_a")
                b = c2.selectbox("al activo", [x for x in assets if x != a], format_func=ticker_label, key="vb_rel_b")
                longs, shorts = [a], [b]
                q_label, q_default = "Diferencial anual esperado (%)", 3.0
            else:
                longs = st.multiselect("Canasta larga (supera)", assets, format_func=ticker_label, key="vb_bl")
                shorts = st.multiselect("Canasta corta (opcional)", [x for x in assets if x not in longs],
                                        format_func=ticker_label, key="vb_bs",
                                        help="Déjala vacía para una view absoluta sobre la canasta.")
                weighting = st.radio("Ponderación dentro de la canasta", ["Igual", "Capitalización"],
                                     horizontal=True, key="vb_w")
                q_label = "Diferencial anual esperado (%)" if shorts else "Rendimiento total anual esperado (%)"
                q_default = 2.0
            q = st.number_input(q_label, min_value=-50.0, max_value=150.0, value=q_default, step=0.5,
                                key=f"vb_q_{vtype}")
            conf = st.slider("Confianza en la view (%)", 5, 100, 50, step=5, key="vb_conf",
                             help="Con Idzorek, fracción del camino entre el equilibrio y tu view que recorre el modelo.")
            margin = 5.0
            if OMEGA_METHODS[s.k_omega] == "interval":
                margin = st.slider("Margen del intervalo ± (%)", 0.5, 25.0, 5.0, step=0.5, key="vb_margin")
            if st.button("Agregar view", type="primary", width="stretch", icon=":material/add:"):
                if not longs:
                    st.error("Selecciona al menos un activo.")
                else:
                    if vtype == "Absoluta":
                        desc = f"{longs[0]} rendirá {q:.1f}% anual"
                    elif vtype == "Relativa":
                        desc = f"{longs[0]} superará a {shorts[0]} por {q:.1f}%"
                    else:
                        desc = f"[{' + '.join(longs)}]" + (f" vs [{' + '.join(shorts)}] por {q:.1f}%" if shorts
                                                            else f" rendirá {q:.1f}% anual")
                    row = pd.DataFrame([[True, vtype, ", ".join(longs), ", ".join(shorts), float(q), float(conf),
                                         weighting, float(q - margin), float(q + margin), desc]], columns=VIEW_COLUMNS)
                    base = s.get("views_current", s.views_df)
                    s.views_df = pd.concat([base[VIEW_COLUMNS], row], ignore_index=True)
                    s.ed_ver += 1
                    st.toast(f"View agregada: {desc}", icon=":material/check_circle:")
                    st.rerun()
    with t_col:
        st.markdown("**Views registradas**")
        vdisp = s.views_df.copy()
        vdisp.insert(0, "Eliminar", False)
        vedit = st.data_editor(
            vdisp, hide_index=True, width="stretch", key=f"ed_views_{ver}",
            column_order=["Activa", "Descripción", "Rendimiento esperado (%)", "Confianza (%)", "Eliminar"],
            disabled=["Descripción", "Tipo"],
            column_config={
                "Activa": st.column_config.CheckboxColumn(width="small", help="Incluir en el modelo"),
                "Descripción": st.column_config.TextColumn(width="medium"),
                "Tipo": st.column_config.TextColumn(width="small"),
                "Rendimiento esperado (%)": st.column_config.NumberColumn("Rend. (%)", format="%.2f", step=0.5,
                                                                          width="small"),
                "Confianza (%)": st.column_config.NumberColumn("Confianza (%)", min_value=1.0, max_value=100.0,
                                                               step=5.0, format="%.0f", width="small"),
                "Eliminar": st.column_config.CheckboxColumn(width="small"),
            })
        views_edit = vedit.drop(columns="Eliminar")[VIEW_COLUMNS]
        s.views_current = views_edit
        c1, c2, c3 = st.columns(3)
        if c1.button("Eliminar marcadas", width="stretch", icon=":material/delete:", disabled=not vedit["Eliminar"].any()):
            s.views_df = views_edit[~vedit["Eliminar"].values].reset_index(drop=True)
            s.ed_ver += 1
            st.rerun()
        if c2.button("Views de ejemplo", width="stretch", icon=":material/auto_awesome:"):
            s.views_df = example_views_df(assets)
            s.ed_ver += 1
            st.rerun()
        if c3.button("Borrar todas", width="stretch", icon=":material/clear_all:"):
            s.views_df = empty_views_df()
            s.ed_ver += 1
            st.rerun()
        if views_edit.empty:
            st.caption("Sin views: el modelo reproduce el portafolio de equilibrio (control de consistencia).")

with t_opt:
    section("Restricciones del portafolio", "Activa las restricciones que quieras aplicar.", "1")
    c1, c2, c3 = st.columns(3)
    tab_widget("toggle", "Límites de peso por activo", "k_use_bounds", c1)
    tab_widget("toggle", "Límites por sector / grupo", "k_use_groups", c2)
    tab_widget("toggle", "Rotación máxima (turnover)", "k_use_turnover", c3)
    if s.k_use_bounds:
        lo_min = 0.0 if s.k_long_only else -50.0
        rng = s.k_wrange
        rng = (max(float(rng[0]), lo_min), float(rng[1]))
        s.k_wrange = rng
        tab_widget("slider", "Rango de peso permitido para cada activo (%)", "k_wrange", min_value=lo_min,
                   max_value=100.0, step=1.0, help="Se aplica a todos los activos, salvo que personalices los límites.")
        tab_widget("toggle", "Personalizar límites por activo", "k_custom_bounds")
    if s.k_use_turnover:
        tab_widget("slider", "Rotación máxima ‖w − w_actual‖₁ (%)", "k_turnover", min_value=0.0, max_value=200.0,
                   step=5.0, help="Si no capturas el portafolio actual, la rotación se mide contra w_mkt.")
    cols = ["Ticker", "Nombre"]
    if s.k_use_groups:
        cols.append("Grupo")
    if s.k_use_bounds and s.k_custom_bounds:
        cols += ["Mín (%)", "Máx (%)"]
    if s.k_use_turnover:
        cols.append("Peso actual (%)")
    if len(cols) > 2:
        ed = st.data_editor(s.assets_df[cols], hide_index=True, disabled=["Ticker", "Nombre"],
                            key=f"ed_assets_{ver}_{'-'.join(cols)}", width="stretch",
                            column_config={
                                "Grupo": st.column_config.SelectboxColumn(options=GROUPS, required=True),
                                "Mín (%)": st.column_config.NumberColumn(min_value=-100.0, max_value=100.0, step=1.0),
                                "Máx (%)": st.column_config.NumberColumn(min_value=0.0, max_value=200.0, step=1.0),
                                "Peso actual (%)": st.column_config.NumberColumn(min_value=-100.0, max_value=200.0,
                                                                                 step=0.5, help="Portafolio actual"),
                            })
        for ccol in cols[2:]:
            s.assets_df[ccol] = ed[ccol].values
    if s.k_use_groups:
        groups_present = sorted({str(g) for g in s.assets_df["Grupo"].fillna("") if str(g).strip()})
        gdf = s.groups_df
        gdf = pd.concat([gdf, pd.DataFrame({"Grupo": [g for g in groups_present if g not in set(gdf["Grupo"])],
                                            "Mín (%)": 0.0, "Máx (%)": 100.0})], ignore_index=True)
        gdf = gdf[gdf["Grupo"].isin(groups_present)].reset_index(drop=True)
        st.markdown("**Límites por sector / grupo**")
        gedit = st.data_editor(gdf, hide_index=True, disabled=["Grupo"], key=f"ed_groups_{ver}_{len(gdf)}",
                               width="stretch", column_config={
                                   "Mín (%)": st.column_config.NumberColumn(min_value=0.0, max_value=100.0, step=5.0),
                                   "Máx (%)": st.column_config.NumberColumn(min_value=0.0, max_value=100.0, step=5.0)})
        s.groups_df = gedit

# =============================================================================
# Configuración y ejecución del modelo
# =============================================================================
adf = s.assets_df.set_index("Ticker")
wlo, whi = s.k_wrange
cfg = ModelConfig(
    tickers=assets, start=s.data_params[1], end=s.data_params[2], benchmark=s.data_params[3],
    max_missing_pct=s.data_params[4], rf_method=RF_METHODS[s.k_rf_method], rf_manual_pct=s.k_rf_manual,
    wmkt_method=WMKT_METHODS[s.k_wmkt],
    manual_weights={t: float(adf.loc[t, "Peso manual w_mkt (%)"] or 0) / 100 for t in assets},
    delta_method=DELTA_METHODS[s.k_delta], delta_manual=s.k_delta_manual, delta_bounds=(s.k_delta_lo, s.k_delta_hi),
    cov_method=COV_METHODS[s.k_cov], ewma_lambda=s.k_ewma, tau_method=TAU_METHODS[s.k_tau],
    tau_manual=s.k_tau_manual, omega_method=OMEGA_METHODS[s.k_omega], interval_prob_pct=s.k_int_prob,
    sigma_opt=SIGMA_OPT[s.k_sigma_opt], budget=BUDGET_MODES[s.k_budget], long_only=s.k_long_only,
    use_asset_bounds=s.k_use_bounds, default_min_pct=wlo, default_max_pct=whi,
    asset_bounds=({t: [float(adf.loc[t, "Mín (%)"]), float(adf.loc[t, "Máx (%)"])] for t in assets}
                  if s.k_custom_bounds else {t: [float(wlo), float(whi)] for t in assets}),
    use_groups=s.k_use_groups,
    asset_groups={t: str(adf.loc[t, "Grupo"]) for t in assets if str(adf.loc[t, "Grupo"]).strip()},
    group_bounds={r_["Grupo"]: [float(r_["Mín (%)"]) if pd.notna(r_["Mín (%)"]) else None,
                                float(r_["Máx (%)"]) if pd.notna(r_["Máx (%)"]) else None]
                  for _, r_ in s.groups_df.iterrows()},
    use_turnover=s.k_use_turnover, turnover_max_pct=s.k_turnover,
    current_weights={t: float(adf.loc[t, "Peso actual (%)"] or 0) for t in assets},
    views=views_to_records(views_edit),
)
res = run_model(md, cfg, views_edit)
rf = res.rf
mu_tot = res.post.mu_bl + rf
ante = ex_ante_summary(res.w_opt, res.post.mu_bl, res.sigma_opt, rf) if res.opt.ok else {}

# Tarjetas de resumen (arriba de las pestañas)
if res.opt.ok:
    mkt_r, mkt_v = portfolio_point(res.w_mkt, res.post.mu_bl, res.sigma_opt)
    kpis([("Rend. esperado", PCT.format(ante["Rendimiento esperado"]),
           f"mercado: {mkt_r + rf:.2%}", "pos" if ante["Rendimiento esperado"] >= mkt_r + rf else "neg"),
          ("Volatilidad", PCT.format(ante["Volatilidad"]), f"mercado: {mkt_v:.2%}", ""),
          ("Sharpe", f"{ante['Sharpe']:.2f}", f"mercado: {mkt_r / mkt_v:.2f}" if mkt_v > 0 else "", ""),
          ("VaR 95 % 1 día", PCT.format(ante["VaR 95% 1 día"]), "paramétrico", "neu"),
          ("Views activas", f"{res.views.k}", f"Ω: {short(s.k_omega)}", "neu"),
          ("Parámetros δ · τ", f"{res.delta:.2f} · {res.tau:.3f}", f"r_f {rf:.2%}", "neu")], container=kpi_box)
else:
    kpi_box.error("El optimizador no encontró solución con las restricciones actuales. Revisa que la suma de "
                  "mínimos sea ≤ 100 % ≤ suma de máximos.")

# ---------------------------------------------------------------------------- Equilibrio
with t_eq:
    for w in [w for w in res.warnings if "δ" in w or "capitalización" in w.lower() or "manual" in w.lower()]:
        st.warning(w)
    section("Rendimientos implícitos de equilibrio", "Π es el rendimiento que haría óptimo al portafolio de "
            "referencia. A diferencia del promedio histórico, es estable y consistente con el riesgo de cada activo.")
    kpis([("δ aversión al riesgo", f"{res.delta:.3f}",
           f"implícito sin acotar {res.delta_raw:.2f}" if res.delta_raw is not None else short(s.k_delta), ""),
          ("τ", f"{res.tau:.4f}", short(s.k_tau), ""),
          ("r_f anual", PCT.format(rf), short(s.k_rf_method), ""),
          ("w_mkt", {"market_cap": "Capitalización", "equal": "1/N", "manual": "Manual"}[res.wmkt_method_used],
           "pesos de referencia", ""),
          ("Estimador Σ", {"ledoit_wolf": "Ledoit-Wolf", "sample": "Muestral", "ewma": "EWMA"}[res.cov.method],
           f"shrinkage {res.cov.shrinkage:.3f}" if res.cov.shrinkage is not None else "", "")])
    eq = pd.DataFrame({"w_mkt": res.w_mkt, "Π exceso": res.pi, "Π total (Π + r_f)": res.pi + rf,
                       "Histórico (aritm.)": res.hist["Rend. aritmético anual"],
                       "Volatilidad": np.sqrt(np.diag(res.cov.cov.values))})
    l, r = st.columns([1, 1.3])
    l.dataframe(fmt_pct_df(eq), width="stretch")
    r.plotly_chart(grouped_bars(eq[["Π total (Π + r_f)", "Histórico (aritm.)"]],
                                       "Equilibrio vs histórico"), width="stretch")
    section("Diagnóstico de la matriz de covarianza")
    kpis([("Número de condición", f"{res.cov.condition_number:,.1f}",
           "> 1,000 indica matriz mal condicionada", "pos" if res.cov.condition_number < 1000 else "neg"),
          ("Eigenvalor mínimo", f"{res.cov.min_eigenvalue:.2e}", "debe ser > 0", "neu"),
          ("Reparación PSD", "Sí" if res.cov.psd_repaired else "No", "semidefinida positiva", "neu")])
    st.plotly_chart(heatmap(cov_to_corr(res.cov.cov)), width="stretch")

# ---------------------------------------------------------------------------- Views
with t_views:
    for e in res.views.errors:
        st.error(e)
    for w in res.views.warnings:
        st.warning(w)
    section("Impacto de las views", None, "2")
    if res.views.k == 0:
        st.info("Sin views activas: μ_BL = Π y el portafolio óptimo sin restricciones reproduce w_mkt.")
    else:
        vt = pd.DataFrame({"Q capturado": res.views.Q_input, "Q exceso r_f": res.views.Q,
                           "Equilibrio PΠ": res.views.P.values @ res.pi.values,
                           "Sorpresa Q − PΠ": res.post.view_gap, "Posterior Pμ_BL": res.post.view_posterior,
                           "Confianza capturada": res.views.confidence, "Confianza implícita": res.implied_conf,
                           "ω": np.diag(res.omega)})
        st.dataframe(vt.style.format({**{c: PCT for c in vt.columns if c != "ω"}, "ω": "{:.6f}"}), width="stretch")
        st.markdown('<div class="bl-note">Confianza implícita = pτΣpᵀ / (pτΣpᵀ + ω): fracción del camino entre el '
                    'equilibrio y la view que recorre el modelo. Con Idzorek coincide con la confianza capturada.</div>',
                    unsafe_allow_html=True)
        post_tbl = pd.DataFrame({"Π total": res.pi + rf, "μ_BL total": mu_tot})
        l, r = st.columns(2)
        l.plotly_chart(grouped_bars(post_tbl, "Equilibrio vs posterior"), width="stretch")
        r.plotly_chart(stacked_contrib(res.post.view_contrib, "Aporte de cada view a μ_BL − Π"),
                       width="stretch")
        st.plotly_chart(stacked_contrib(res.w_unconstrained_contrib,
                                               "Aporte de cada view a los pesos (sin restricciones)"), width="stretch")
        with st.expander("Matriz P (selección de activos por view)"):
            st.dataframe(res.views.P.style.format("{:.3f}"), width="stretch")

# ---------------------------------------------------------------------------- Optimización
fr_post = cached_frontier(_hash(res.post.mu_bl, res.sigma_opt, res.constraints), res.post.mu_bl, res.sigma_opt,
                          res.constraints)
with t_opt:
    for e in [e for e in res.errors if e not in res.views.errors]:
        st.error(e)
    section("Portafolio óptimo", f"Estado del optimizador: {res.opt.status}", "2")
    wt = pd.DataFrame({"w_mkt (equilibrio)": res.w_mkt, "w* sin restricciones": res.w_unconstrained,
                       "w* óptimo": res.w_opt, "Activo vs mercado": res.w_opt - res.w_mkt, "μ_BL total": mu_tot})
    l, r = st.columns([1, 1.3])
    l.dataframe(fmt_pct_df(wt), width="stretch")
    l.markdown(f'<div class="bl-note">Suma w* óptimo: {res.w_opt.sum():.2%} · '
               f'suma w* sin restricciones: {res.w_unconstrained.sum():.2%}</div>', unsafe_allow_html=True)
    r.plotly_chart(grouped_bars(wt[["w_mkt (equilibrio)", "w* óptimo"]], "Pesos: equilibrio vs óptimo"),
                   width="stretch")
    ok = res.consistency_gap < 1e-6
    st.markdown(f'<div class="bl-note">{"✅" if ok else "⚠️"} <b>Control de consistencia:</b> sin views y sin '
                f'restricciones, w* normalizado = w_mkt (desviación máxima {res.consistency_gap:.1e}).</div>',
                unsafe_allow_html=True)

    section("Fronteras eficientes", "Portafolios totalmente invertidos con las mismas restricciones de pesos.", "3")
    show_all = st.toggle("Comparar con fronteras histórica y de equilibrio", value=True)
    curves = {"Posterior (μ_BL)": fr_post.assign(ret=fr_post["ret"] + rf) if not fr_post.empty else fr_post}
    if show_all:
        mu_hist_ex = res.hist["Rend. aritmético anual"] - rf
        fh = cached_frontier(_hash(mu_hist_ex, res.cov.cov, res.constraints), mu_hist_ex, res.cov.cov, res.constraints)
        fe = cached_frontier(_hash(res.pi, res.cov.cov, res.constraints), res.pi, res.cov.cov, res.constraints)
        curves["Histórica (μ̂)"] = fh.assign(ret=fh["ret"] + rf) if not fh.empty else fh
        curves["Equilibrio (Π)"] = fe.assign(ret=fe["ret"] + rf) if not fe.empty else fe
    pts = {}
    for name, w in (("Óptimo BL", res.w_opt), ("Mercado", res.w_mkt),
                    ("1/N", pd.Series(1 / len(assets), index=assets))):
        if w.notna().all():
            rr, vv = portfolio_point(w, res.post.mu_bl, res.sigma_opt)
            pts[name] = (rr + rf, vv)
    asset_pts = {a: (mu_tot[a], float(np.sqrt(res.sigma_opt.loc[a, a]))) for a in assets}
    st.plotly_chart(frontiers(curves, pts, asset_pts), width="stretch")

# ---------------------------------------------------------------------------- Riesgo
with t_risk:
    if not res.opt.ok:
        st.error("El optimizador no encontró solución; ajusta las restricciones.")
    else:
        section("Riesgo ex-ante", "Con μ_BL y la covarianza usada en la optimización.")
        kpis([("VaR 95 % 1 día", PCT.format(ante["VaR 95% 1 día"]), "pérdida no superada el 95 % de los días", "neg"),
              ("CVaR 95 % 1 día", PCT.format(ante["CVaR 95% 1 día"]), "pérdida media en el 5 % peor", "neg"),
              ("VaR 95 % 1 año", PCT.format(ante["VaR 95% 1 año"]), "horizonte anual", "neg"),
              ("Diversificación", f"{ante['Ratio de diversificación']:.2f}", "Σ wᵢσᵢ / σₚ", "pos"),
              ("N efectivo", f"{ante['N efectivo (1/Σw²)']:.1f}", "activos equivalentes", "neu")])
        rc = risk_contributions(res.w_opt, res.sigma_opt)
        l, r = st.columns([1, 1.3])
        l.dataframe(fmt_pct_df(rc), width="stretch")
        r.plotly_chart(grouped_bars(rc[["Peso", "Contribución (%)"]], "Peso vs contribución al riesgo"),
                       width="stretch")
        st.markdown('<div class="bl-note">Contribución de Euler: RCᵢ = wᵢ(Σw)ᵢ/σₚ; la suma es la volatilidad total.'
                    '</div>', unsafe_allow_html=True)
        section("Riesgo ex-post", "Pesos óptimos aplicados a la historia (descriptivo, dentro de muestra).")
        post = ex_post_summary(md.returns, res.w_opt, md.rf_daily, md.benchmark_returns)
        ratio_keys = ("Sharpe", "Sortino", "Calmar", "Beta", "Information", "Asimetría", "Curtosis")
        post_df = pd.DataFrame({"Valor": [f"{v:.3f}" if any(k in n for k in ratio_keys) else PCT.format(v)
                                          for n, v in post.items()]}, index=list(post))
        l, r = st.columns([1, 1.3])
        l.dataframe(post_df, width="stretch")
        pr = portfolio_returns(md.returns, res.w_opt, md.rf_daily)
        comp = pd.DataFrame({"Óptimo BL": pr, "Índice de referencia": md.benchmark_returns.reindex(pr.index)})
        r.plotly_chart(drawdown_curves(comp, "Drawdown histórico"), width="stretch")

# ---------------------------------------------------------------------------- Backtest
with t_bt:
    section("Backtest walk-forward", "En cada rebalanceo se reestiman Σ, δ y Π solo con datos previos y se optimiza "
            "con las mismas restricciones. Compara BL con views, equilibrio sin views, 1/N y el índice.")
    with st.container(border=True):
        c = st.columns([1.2, 1.2, 1, 1])
        freq = c[0].segmented_control("Rebalanceo", ["Mensual", "Trimestral"], default="Mensual", required=True,
                                      key="bt_freq")
        lookback = c[1].select_slider("Ventana de estimación", options=[126, 252, 504, 756], value=252,
                                      format_func=lambda d_: {126: "6 meses", 252: "1 año", 504: "2 años",
                                                              756: "3 años"}[d_], key="bt_lb")
        cost = c[2].number_input("Costo (pb por operación)", 0.0, 200.0, 10.0, step=1.0, key="bt_cost")
        c[3].markdown("<div style='height:1.7rem'></div>", unsafe_allow_html=True)
        run_bt = c[3].button("Ejecutar", type="primary", width="stretch", icon=":material/play_arrow:")
    if run_bt:
        bar = st.progress(0.0, text="Ejecutando backtest…")
        try:
            vdf = views_edit if res.views.k else None
            s.backtest = run_backtest(md, cfg, vdf, "M" if freq == "Mensual" else "Q", int(lookback), float(cost),
                                      progress=lambda p_: bar.progress(p_, text=f"Rebalanceos: {p_:.0%}"))
            s.exports = None
        except Exception as e:  # noqa: BLE001
            st.error(str(e))
        bar.empty()
    bt = s.backtest
    if bt is not None:
        for n in bt.notes:
            st.warning(n)
        main = "BL con views" if "BL con views" in bt.stats.columns else "Equilibrio (sin views)"
        st_ = bt.stats[main]
        kpis([("CAGR", PCT.format(st_["CAGR"]), main, "pos" if st_["CAGR"] >= 0 else "neg"),
              ("Volatilidad", PCT.format(st_["Volatilidad anual"]), "anual", ""),
              ("Sharpe", f"{st_['Sharpe']:.2f}", f"índice: {bt.stats['Índice de referencia']['Sharpe']:.2f}", ""),
              ("Máx. drawdown", PCT.format(st_["Máx. drawdown"]), "peor caída", "neg"),
              ("Tracking error", PCT.format(st_["Tracking error"]), "vs índice", "neu")])
        st.plotly_chart(value_curves(bt.returns), width="stretch")
        st.plotly_chart(drawdown_curves(bt.returns), width="stretch")
        ratio_rows = [x for x in ("Sharpe", "Sortino", "Calmar", "Beta", "Information ratio") if x in bt.stats.index]
        st.dataframe(bt.stats.style.format(lambda v: "—" if pd.isna(v) else f"{v:.2%}")
                     .format(lambda v: "—" if pd.isna(v) else f"{v:.3f}", subset=pd.IndexSlice[ratio_rows, :]),
                     width="stretch")
        st.plotly_chart(weights_area(bt.weights[main], f"Pesos objetivo — {main}"), width="stretch")
    else:
        st.info("Configura los parámetros y pulsa **Ejecutar**.")

# ---------------------------------------------------------------------------- Exportar
with t_exp:
    section("Exportar resultados", "Excel con todas las tablas, matrices y parámetros; reporte PDF con metodología, "
            "resultados y gráficas.")
    if not res.opt.ok:
        st.error("No hay portafolio óptimo que exportar; ajusta las restricciones.")
    else:
        with st.container(border=True):
            if st.button("Generar archivos", type="primary", icon=":material/description:"):
                with st.spinner("Generando Excel y PDF…"):
                    post = ex_post_summary(md.returns, res.w_opt, md.rf_daily, md.benchmark_returns)
                    tables = results_tables(res, md, ante, post, fr_post, s.backtest)
                    s.exports = {"xlsx": to_excel(tables), "pdf": to_pdf(res, md, ante, fr_post, s.backtest)}
            if s.exports:
                stamp = date.today().strftime("%Y%m%d")
                c1, c2 = st.columns(2)
                c1.download_button("Descargar Excel", s.exports["xlsx"], f"black_litterman_{stamp}.xlsx",
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                   width="stretch", icon=":material/table_view:")
                c2.download_button("Descargar reporte PDF", s.exports["pdf"], f"black_litterman_{stamp}.pdf",
                                   "application/pdf", width="stretch", icon=":material/picture_as_pdf:")
            if s.backtest is None:
                st.caption("Sugerencia: ejecuta el backtest antes de generar los archivos para incluirlo en el reporte.")

st.markdown('<div class="bl-note" style="margin-top:28px;text-align:center">Herramienta analítica con fines '
            'educativos y de investigación; no constituye asesoría de inversión. Los datos de Yahoo Finance pueden '
            'contener errores o retrasos.</div>', unsafe_allow_html=True)
